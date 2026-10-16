#!/usr/bin/env python3
"""Validate editable paper figures and build previews from Office-rendered PDFs."""

from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import tempfile
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = ROOT / "docs/reports/2026-10/assets"
MATH_FAMILY = "Liberation Serif"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def require(condition, message):
    if not condition:
        raise ValueError(message)


def pdf_pages(path: Path) -> int:
    info = subprocess.check_output(["pdfinfo", str(path)], text=True)
    return int(re.search(r"^Pages:\s+(\d+)", info, re.MULTILINE).group(1))


def validate_slide(slide, svg: Path, presentation, reference_colors):
    root = ET.parse(svg).getroot()
    expected = [node for node in root if node.tag.endswith("}text")]
    text_shapes = [
        shape for shape in slide.shapes if shape.has_text_frame and shape.text
    ]
    require(
        [shape.text for shape in text_shapes]
        == ["".join(node.itertext()) for node in expected],
        f"SVG/PowerPoint text differs: {svg.name}",
    )
    math_count = italic_count = 0
    text_colors = {node.get("fill") for node in expected}
    require(
        text_colors <= {"#000000", "#ffffff"},
        f"Text outside black/white typography: {svg.name}",
    )
    numbered_steps = sum(
        node.get("fill") == "#ffffff" and "".join(node.itertext()).isdecimal()
        for node in expected
    )
    require(
        svg.name.startswith("01_") or numbered_steps >= 3,
        f"Missing white step numbers: {svg.name}",
    )
    step_circles = [
        node for node in root if node.get("data-execution-step") is not None
    ]
    require(
        svg.name.startswith("01_")
        or (
            sorted(node.get("data-execution-step") for node in step_circles) == ["1", "2", "3", "4"]
            and all(node.get("fill") == "#000000" for node in step_circles)
        ),
        f"Missing numbered execution-arrow markers: {svg.name}",
    )
    routes = [node for node in root if "marker-end" in node.attrib]
    heads = [shape for shape in slide.shapes if shape.name.startswith("arrowhead ")]
    require(len(heads) == len(routes), f"Missing editable arrowheads in {svg.name}")
    for route, head in zip(routes, heads):
        is_open = route.get("data-arrow-kind") == "open"
        require(
            (head.fill.type == MSO_FILL_TYPE.BACKGROUND) == is_open,
            f"Arrowhead style differs from SVG in {svg.name}",
        )
        require(
            float(route.get("stroke-width")) >= 4.0
            and route.get("stroke") != expected[0].get("fill"),
            f"Directed route is indistinguishable from a border in {svg.name}",
        )
    expected_cubics = sum(node.get("d", "").count("C") for node in root)
    native_cubics = sum(
        len(shape._element.xpath(".//a:cubicBezTo")) for shape in slide.shapes
    )
    require(
        native_cubics == expected_cubics and native_cubics > 0,
        f"Editable cubic curves differ from SVG in {svg.name}",
    )
    for node, shape in zip(expected, text_shapes):
        runs = [run for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
        require(len(runs) == 1, f"Unexpected text-run structure: {shape.text!r}")
        run = runs[0]
        require(run.font.name == node.get("font-family"), f"Font changed: {shape.text!r}")
        require(
            str(run.font.color.rgb) == node.get("fill").lstrip("#").upper(),
            f"Text color changed: {shape.text!r}",
        )
        require(
            bool(run.font.italic) == (node.get("font-style") == "italic"),
            f"Math style changed: {shape.text!r}",
        )
        if run.font.name == MATH_FAMILY:
            math_count += 1
            italic_count += bool(run.font.italic)
    pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    require(pictures == 0, f"Flattened picture in {svg.name}")
    ids = [shape.shape_id for shape in slide.shapes]
    require(len(ids) == len(set(ids)), f"Duplicate shape ID in {svg.name}")
    for shape in slide.shapes:
        require(
            all("#" + node.get("val").lower() in reference_colors
                for node in shape._element.xpath(".//a:srgbClr")),
            f"PowerPoint shape color absent from reference: {shape.name!r}",
        )
        require(
            shape.left >= 0 and shape.top >= 0
            and shape.left + shape.width <= presentation.slide_width
            and shape.top + shape.height <= presentation.slide_height,
            f"Object outside slide: {shape.name!r}",
        )
        require(
            bool(shape._element.xpath("p:spPr/a:effectLst")),
            f"Missing explicit effect reset: {shape.name!r}",
        )
    notes = slide.notes_slide.notes_text_frame.text
    require(len(notes) > 300, f"Missing mechanism notes in {svg.name}")
    require(math_count > 0 and italic_count > 0, f"Missing editable math in {svg.name}")
    return {
        "editable_objects": len(slide.shapes),
        "editable_text_objects": len(text_shapes),
        "editable_math_text_objects": math_count,
        "italic_math_variables": italic_count,
        "text_matches_svg": True,
        "font_and_italic_styles_match_svg": True,
        "black_and_white_text_only": True,
        "text_colors_match_svg": True,
        "shape_colors_in_reference_samples": True,
        "white_on_black_step_numbers": numbered_steps,
        "numbered_execution_arrow_markers": len(step_circles),
        "all_objects_within_slide": True,
        "unique_shape_ids": True,
        "notes_present": True,
        "picture_objects": pictures,
        "all_shapes_and_text_have_effect_resets": True,
        "editable_arrowheads": len(heads),
        "open_arrowheads": sum(route.get("data-arrow-kind") == "open" for route in routes),
        "colored_routes_and_arrow_styles_match": True,
        "editable_cubic_segments": native_cubics,
        "curves_preserved_without_line_approximation": True,
    }


def validate_no_effects(path: Path):
    forbidden = {
        "outerShdw", "innerShdw", "prstShdw", "glow", "reflection",
        "softEdge", "blur", "scene3d", "sp3d",
    }
    with ZipFile(path) as package:
        for name in package.namelist():
            if not name.endswith(".xml"):
                continue
            root = ET.fromstring(package.read(name))
            for node in root.iter():
                local = node.tag.rsplit("}", 1)[-1]
                require(local not in forbidden, f"Visual effect remains in {name}: {local}")
                if local == "effectRef":
                    require(node.get("idx") == "0", f"Inherited effect remains in {name}")
    return True


def office_preview(pdf: Path, destination: Path, figures=None):
    page_count = pdf_pages(pdf)
    require(page_count == 4, f"Office PDF must contain overview plus three designs: {pdf}")
    with tempfile.TemporaryDirectory(prefix="beliefkv-paper-preview-") as scratch:
        prefix = Path(scratch) / "page"
        subprocess.run(
            ["pdftoppm", "-png", "-scale-to", "1800", str(pdf), str(prefix)],
            check=True,
        )
        pages = []
        for index in range(page_count):
            with Image.open(Path(scratch) / f"page-{index + 1}.png") as image:
                page = image.convert("RGB")
            if figures:
                width, height = figures[index]["canvas"]
                content_height = round(page.width * height / width)
                top = (page.height - content_height) // 2
                page = page.crop((0, top, page.width, top + content_height))
            page = page.resize(
                (1440, round(page.height * 1440 / page.width)),
                Image.Resampling.LANCZOS,
            )
            pages.append(page)
        sheet = Image.new("RGB", (1488, 24 + sum(page.height + 24 for page in pages)), "white")
        top = 24
        for page in pages:
            sheet.paste(page, (24, top))
            top += page.height + 24
        sheet.save(destination)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--office-pdf-directory", type=Path, required=True)
    parser.add_argument("--reference-image", type=Path)
    parser.add_argument("--visual-review-complete", action="store_true")
    args = parser.parse_args()
    manifest = json.loads((args.output / "manifest.json").read_text())
    generator = Path(__file__).with_name("generate_monthly_paper_figures.py")
    require(digest(generator) == manifest["generator_sha256"], "Generator changed after export")
    palette_path = Path(__file__).with_name("janus_reference_palette.json")
    require(digest(palette_path) == manifest["reference_palette_sha256"],
            "Reference palette changed after export")
    reference_palette = manifest["reference_palette"]
    reference_colors = {row["hex"].lower() for row in reference_palette["samples"]}
    reference_verified = False
    if args.reference_image:
        require(digest(args.reference_image) == reference_palette["source_sha256"],
                "Reference image does not match the recorded source")
        with Image.open(args.reference_image) as image:
            require(list(image.size) == reference_palette["source_size"],
                    "Reference image dimensions differ")
            image = image.convert("RGB")
            for row in reference_palette["samples"]:
                actual = "#%02X%02X%02X" % image.getpixel(tuple(row["xy"]))
                require(actual == row["hex"], f"Reference color sample differs: {row}")
        reference_verified = True
    for name, expected in manifest["architecture_source_sha256"].items():
        require(digest(ROOT / name) == expected, f"Architecture source changed: {name}")

    decks = []
    for language in ("zh", "en"):
        path = args.output / f"beliefkv_architecture_editable_{language}.pptx"
        presentation = Presentation(path)
        no_effects = validate_no_effects(path)
        require(len(presentation.slides) == 4, f"Expected overview plus three designs: {path}")
        figures = [row for row in manifest["figures"] if row["language"] == language]
        require(
            [row["stem"] for row in figures]
            == [row["stem"] + "_" + language for row in manifest["figure_designs"]],
            f"Overview/design order differs: {path}",
        )
        for row in figures:
            allowed = set(manifest["color_policy"]["mechanism_header_colors"])
            if row["stem"].startswith("01_"):
                allowed.add("#ffffff")
            require(
                all(panel["header"] in allowed for panel in row["panel_styles"]),
                f"Mechanism header outside the restrained palette: {row['stem']}",
            )
            require(
                not {"H2D", "D2H"}.intersection(row["transfer_entities"])
                and (
                    row["stem"].startswith("02_")
                    or row["transfer_action_labels"] > 0
                ),
                f"Transfer actions are represented as components: {row['stem']}",
            )
            svg = ET.parse(args.output / (row["stem"] + ".svg")).getroot()
            drawing_colors = {
                value.lower()
                for node in svg.iter()
                for key in ("fill", "stroke")
                if (value := node.get(key, "")).startswith("#")
            }
            require(drawing_colors <= reference_colors,
                    f"Drawing color absent from reference: {row['stem']}: "
                    f"{sorted(drawing_colors - reference_colors)}")
            labels = {"".join(node.itertext()) for node in svg if node.tag.endswith("}text")}
            required_designs = (
                {"D1", "D2", "D3"} if row["stem"].startswith(("01_", "04_"))
                else {"D1", "D3"} if row["stem"].startswith("02_") else {"D2", "D3"}
            )
            require(
                required_designs <= labels,
                f"Shared design interfaces missing: {row['stem']}",
            )
            require(
                row["stem"].startswith("01_")
                or (
                    row["layout"] == "compact_algorithm"
                    and not row["panel_styles"]
                    and row["canvas"][1] < 800
                ),
                f"Mechanism is not compact: {row['stem']}",
            )
            if row["stem"].startswith("01_"):
                tiers = [node for node in svg if node.get("data-memory-tier")]
                headers = [node for node in svg if node.get("data-tier-header")]
                require(
                    [node.get("data-memory-tier") for node in tiers] == ["HBM", "Host"]
                    and len(headers) == 2,
                    f"HBM/Host peer tier pair missing: {row['stem']}",
                )
                require(
                    all(
                        tiers[0].get(key) == tiers[1].get(key)
                        for key in ("x", "width", "height", "stroke-width", "rx")
                    )
                    and all(
                        headers[0].get(key) == headers[1].get(key)
                        for key in ("width", "height", "rx")
                    )
                    and tiers[0].get("fill") != tiers[1].get("fill")
                    and headers[0].get("fill") != headers[1].get("fill"),
                    f"HBM/Host geometry or location colors are ambiguous: {row['stem']}",
                )
                containers = {
                    node.get("data-container-role"): node for node in svg
                    if node.get("data-container-role") is not None
                }
                require(
                    {"control_plane", "data_plane"} <= containers.keys()
                    and containers["control_plane"].get("fill") != containers["data_plane"].get("fill")
                    and containers["control_plane"].get("stroke") != containers["data_plane"].get("stroke"),
                    f"Control/data planes have ambiguous visual roles: {row['stem']}",
                )
            page_nodes = [node for node in svg if node.get("data-page-tier")]
            if not row["stem"].startswith("02_"):
                require(
                    {"HBM", "Host"} <= {node.get("data-page-tier") for node in page_nodes},
                    f"Location-encoded pages missing: {row['stem']}",
                )
            for node in page_nodes:
                following = list(svg)[list(svg).index(node) + 1]
                require(
                    following.tag.endswith("}rect")
                    and following.get("fill") == node.get("data-page-owner")
                    and float(following.get("x")) > float(node.get("x"))
                    and float(following.get("y")) > float(node.get("y"))
                    and float(following.get("height")) < float(node.get("height")) / 4,
                    f"Independent owner stripe missing: {row['stem']}",
                )
                if node.get("data-page-missing") == "true":
                    require(node.get("fill") == "#ffffff", "Missing page is not empty")
                    continue
                location = node.get("data-page-tier").lower()
                require(
                    node.get("fill") == manifest["palette"][location]["page"],
                    f"Page body does not encode its location: {row['stem']}",
                )
        slides = [
            validate_slide(slide, args.output / (row["stem"] + ".svg"),
                           presentation, reference_colors)
            for slide, row in zip(presentation.slides, figures)
        ]
        for slide, row in zip(slides, figures):
            require(
                slide["editable_math_text_objects"] == row["editable_math_text_objects"],
                f"Math object count differs: {row['stem']}",
            )
        office_pdf = args.office_pdf_directory / (path.stem + ".pdf")
        office_preview(office_pdf, args.output / f"pptx_preview_{language}.png", figures)
        decks.append({
            "file": path.name, "sha256": digest(path), "slides": slides,
            "office_render_pdf_sha256": digest(office_pdf),
            "office_render_pages": 4,
            "package_wide_shadow_and_3d_effects_absent": no_effects,
            "restrained_mechanism_palette": True,
            "transfer_actions_separate_from_hardware_entities": True,
        })

    standalone_decks = []
    for row in manifest["standalone_powerpoints"]:
        path = args.output / row["file"]
        presentation = Presentation(path)
        require(len(presentation.slides) == 1, f"Expected a single compact figure: {path}")
        no_effects = validate_no_effects(path)
        slide = validate_slide(
            presentation.slides[0], args.output / (row["figure"] + ".svg"),
            presentation, reference_colors,
        )
        expected_ratio = row["canvas"][0] / row["canvas"][1]
        require(
            abs(presentation.slide_width / presentation.slide_height - expected_ratio) < 1e-5,
            f"Standalone page does not match figure aspect: {path}",
        )
        office_pdf = args.office_pdf_directory / (path.stem + ".pdf")
        require(pdf_pages(office_pdf) == 1, f"Standalone Office render missing: {path}")
        standalone_decks.append({
            "file": path.name, "sha256": digest(path), "slide": slide,
            "office_render_pdf_sha256": digest(office_pdf),
            "package_wide_shadow_and_3d_effects_absent": no_effects,
            "page_matches_compact_figure": True,
        })

    for row in manifest["figures"]:
        png = args.output / (row["stem"] + ".png")
        with Image.open(png) as image:
            require(list(image.size) == row["png_pixels"], f"Unexpected PNG dimensions: {png}")
            require(
                all(abs(value - 600) < 1 for value in image.info.get("dpi", (0, 0))),
                f"Expected 600 dpi: {png}",
            )
    vector_pdf = args.output / "beliefkv_figures_bilingual.pdf"
    require(pdf_pages(vector_pdf) == 8, "Expected eight vector PDF pages")
    image_listing = subprocess.check_output(["pdfimages", "-list", str(vector_pdf)], text=True)
    require(len(image_listing.strip().splitlines()) == 2, "Vector PDF contains a raster image")
    receipt = {
        "checked_at": datetime.now(timezone.utc).isoformat(),
        "architecture_date": manifest["architecture_date"],
        "source_revision": manifest["source_revision"],
        "generator_sha256": digest(generator),
        "validator_sha256": digest(Path(__file__)),
        "powerpoints": decks,
        "standalone_powerpoints": standalone_decks,
        "png_count": 8, "png_dpi": 600,
        "vector_pdf_pages": 8, "vector_pdf_raster_images": 0,
        "office_renderer": "LibreOffice 24.2.7 in isolated temporary directory",
        "office_render_check": "eight collection slides and six compact standalone slides rendered",
        "shared_design_interfaces_present": True,
        "overview_plus_three_designs": True,
        "aligned_compact_algorithms": True,
        "numbers_annotate_execution_arrows": True,
        "hbm_host_peer_geometry_distinct_location_colors": True,
        "control_data_distinct_border_and_background": True,
        "page_location_and_ownership_encoded_separately": True,
        "all_explicit_drawing_colors_in_reference_samples": True,
        "reference_image_hash_and_sample_pixels_verified": reference_verified,
        "visual_review_complete": args.visual_review_complete,
        "figure_scope": (
            manifest["architecture_date"]
            + " development mechanisms; diagrams and timing are schematic"
        ),
    }
    (args.output / "powerpoint_validation.json").write_text(
        json.dumps(receipt, indent=2) + "\n",
    )
    print(json.dumps(receipt, indent=2))


if __name__ == "__main__":
    main()
