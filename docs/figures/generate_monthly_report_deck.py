#!/usr/bin/env python3
"""Build the concise monthly report using the supplied PowerPoint template."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


REPORT_DIR = Path(__file__).resolve().parents[1] / "reports" / "2026-10"
TEMPLATE = REPORT_DIR / "月度汇报模板.pptx"
OUTPUT = REPORT_DIR / "BeliefKV_月度汇报_20261011.pptx"
FONT = "微软雅黑"
NAVY = "162A46"
BLUE = "01639C"
TEAL = "357B73"
GRAY = "596875"
RULE = "D6DFE7"
WHITE = "FFFFFF"


def format_paragraph(paragraph, size, color, bold=False, align=PP_ALIGN.LEFT):
    paragraph.alignment = align
    paragraph.line_spacing = 1.12
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    props = paragraph._p.get_or_add_pPr()
    props.set("marL", "0")
    props.set("indent", "0")
    props.append(OxmlElement("a:buNone"))
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        run_props = run._r.get_or_add_rPr()
        run_props.set("lang", "zh-CN")
        run_props.set("spc", "0")
        east_asian = run_props.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}ea"
        )
        if east_asian is None:
            east_asian = OxmlElement("a:ea")
            run_props.append(east_asian)
        east_asian.set("typeface", FONT)


def text(slide, value, x, y, width, height, size=20, color=NAVY,
         bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.text = value
    for paragraph in frame.paragraphs:
        format_paragraph(paragraph, size, color, bold, align)
    return shape


def rectangle(slide, x, y, width, height, fill, stroke=RULE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
        Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(stroke)
    shape.line.width = Pt(1)
    remove_shadow(shape)
    return shape


def remove_shadow(shape):
    shape._element.spPr.append(OxmlElement("a:effectLst"))
    effect_ref = shape._element.find(
        "{http://schemas.openxmlformats.org/presentationml/2006/main}style/"
        "{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef"
    )
    if effect_ref is not None:
        effect_ref.set("idx", "0")


def line(slide, x1, y1, x2, y2, color=RULE, arrow=False, width=1.5):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width)
    remove_shadow(shape)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "sm")
        end.set("len", "sm")
        shape._element.spPr.get_or_add_ln().append(end)
    return shape


def add_slide(presentation, layout, title, notes):
    slide = presentation.slides.add_slide(layout)
    for shape in list(slide.shapes):
        if shape != slide.shapes.title:
            shape._element.getparent().remove(shape._element)
    heading = slide.shapes.title
    heading.text = title
    for paragraph in heading.text_frame.paragraphs:
        format_paragraph(paragraph, 42, WHITE, bold=True)
    text(
        slide, "BeliefKV  |  2026.08.15 - 2026.10.11",
        7.90, 6.80, 4.10, 0.25, size=12, color=GRAY,
        align=PP_ALIGN.RIGHT
    )
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def project_slide(presentation, layout):
    slide = add_slide(
        presentation, layout, "项目介绍",
        "建议用时：55 秒。\n\n"
        "我们的项目叫 BeliefKV，研究动态多 Agent 工作流下的 KV 缓存管理与调度。"
        "一个父 Agent 会创建子任务，也会调用外部工具；等待期间，它的 KV 可能被迁移"
        "到主机内存，恢复执行时又需要搬回 GPU。\n"
        "我们利用子任务收尾通知、输出正文和执行进度，估计子任务或工具何时完成。"
        "预测模块提供完成提示，运行时结合缓存副本、可用空间和执行机会，"
        "决定提前备份、预取，以及恢复后的准入。\n"
        "目标是减少缓存恢复等待，并让恢复好的数据尽快被执行消费，"
        "最终提高工作流吞吐、缩短完成时间。右侧展示了工作流、预测和原生执行之间的关系。"
    )
    text(slide, "BeliefKV", 0.75, 1.78, 4.25, 0.55,
         size=31, color=BLUE, bold=True)
    text(slide, "动态 Agent 工作流的\n预测式 KV 管理与调度",
         0.75, 2.40, 4.3, 0.85, size=22, bold=True)

    for y, label, body in (
        (3.42, "场景", "子任务与工具等待交错，\nKV 在 HBM / Host 间迁移。"),
        (4.45, "方法", "观测收尾信号，提前准备\n下一次执行所需的 KV。"),
        (5.48, "目标", "减少恢复等待，提高吞吐，\n缩短工作流完成时间。"),
    ):
        text(slide, label, 0.75, y, 0.70, 0.4,
             size=21, color=BLUE, bold=True)
        text(slide, body, 1.58, y, 3.64, 0.80, size=20)

    # Distill the existing architecture figure into readable editable modules.
    rectangle(slide, 5.58, 1.90, 7.08, 1.10, "EFF5FA", "ADC5D7")
    text(slide, "动态工作流：子任务与工具调用",
         5.78, 2.06, 6.67, 0.40, size=21, color=BLUE, bold=True,
         align=PP_ALIGN.CENTER)
    for x, width, label in (
        (5.83, 1.25, "父 Agent"),
        (7.61, 1.25, "子 Agent"),
        (9.39, 1.12, "JOIN"),
        (11.05, 1.34, "父 Agent"),
    ):
        text(slide, label, x, 2.53, width, 0.32,
             size=17, align=PP_ALIGN.CENTER)
    for x1, x2 in ((7.13, 7.52), (8.91, 9.30), (10.58, 10.97)):
        line(slide, x1, 2.69, x2, 2.69, BLUE, arrow=True)

    line(slide, 7.25, 3.02, 7.25, 3.42, BLUE, arrow=True)
    rectangle(slide, 5.58, 3.43, 3.34, 1.10, "FBF6EB", "D5BC89")
    text(slide, "完成预测", 5.78, 3.60, 2.94, 0.40,
         size=22, color="926923", bold=True, align=PP_ALIGN.CENTER)
    text(slide, "通知 / 正文 / 执行进度",
         5.74, 4.10, 3.00, 0.28, size=16, color=GRAY,
         align=PP_ALIGN.CENTER)

    line(slide, 8.97, 3.96, 9.29, 3.96, TEAL, arrow=True)
    rectangle(slide, 9.35, 3.43, 3.31, 1.10, "EEF6F3", "ABC9BE")
    text(slide, "迁移与准入决策", 9.50, 3.60, 3.00, 0.40,
         size=22, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "容量 / 缺失页 / 执行机会",
         9.50, 4.10, 3.00, 0.28, size=15.5, color=GRAY,
         align=PP_ALIGN.CENTER)
    line(slide, 11.0, 4.55, 11.0, 4.91, TEAL, arrow=True)

    rectangle(slide, 5.58, 4.93, 7.08, 1.37, WHITE, "ADC5D7")
    text(slide, "SGLang 原生缓存与 GPU 执行",
         5.78, 5.07, 6.67, 0.42, size=22, color=BLUE, bold=True,
         align=PP_ALIGN.CENTER)
    rectangle(slide, 5.89, 5.69, 2.30, 0.45, "EFF5FA", "ADC5D7")
    rectangle(slide, 10.03, 5.69, 2.30, 0.45, "EEF6F3", "ABC9BE")
    text(slide, "GPU / HBM", 5.99, 5.75, 2.10, 0.34,
         size=19, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "CPU / Host", 10.13, 5.75, 2.10, 0.34,
         size=19, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "D2H", 8.51, 5.49, 1.20, 0.22,
         size=12, color=TEAL, align=PP_ALIGN.CENTER)
    text(slide, "H2D", 8.51, 6.11, 1.20, 0.22,
         size=12, color=BLUE, align=PP_ALIGN.CENTER)
    line(slide, 8.31, 5.80, 9.91, 5.80, TEAL, arrow=True)
    line(slide, 9.91, 6.00, 8.31, 6.00, BLUE, arrow=True)
    text(slide, "按缺失页迁移，恢复后及时执行",
         5.78, 6.43, 6.67, 0.32, size=16, color=GRAY,
         align=PP_ALIGN.CENTER)


def progress_slide(presentation, layout):
    slide = add_slide(
        presentation, layout, "上月工作总结",
        "建议用时：75 秒。汇报周期为 2026 年 8 月 15 日至 10 月 11 日。\n\n"
        "这阶段的工作可以分为五段。八月下半月，我们先建立动态工作流的因果图"
        "和调度基础，接通子任务创建、JOIN，以及父任务上下文连续性。\n"
        "九月上中旬，接通预测性主机与 GPU 之间的 KV 迁移，并完善传输确认、"
        "缓存驻留和生命周期。九月下旬，适配 FULL 与 Mamba 混合缓存，"
        "完成轨迹采集、模型训练和项目隔离校准。\n"
        "随后发现，直接预测返回的墙钟时间会混入排队等待，"
        "因此转向结合收尾通知、正文语义和 decode 进度，分别估计执行阶段与剩余工作。\n"
        "十月以来，我们重点做原生 SGLang 对照，并优化增量备份、缺失页恢复、"
        "恢复后的执行准入，以及子任务结束后的容量交接。\n"
        "目前机制链路已贯通，但稳定的端到端性能优势仍需要进一步验证。"
    )
    milestones = (
        ("08.15 - 08.31", "建立工作流与调度基础",
         "因果图、原生 subagent、JOIN 后 parent 上下文连续性。"),
        ("09.01 - 09.21", "贯通预测式 KV 迁移",
         "接通 D2H / H2D、传输确认、驻留与生命周期管理。"),
        ("09.22 - 09.27", "完成混合缓存适配与离线建模",
         "适配 FULL / Mamba；完成 trace 采集、训练与项目隔离校准。"),
        ("09.28 - 10.04", "转向事件与内容驱动的预测",
         "结合通知、正文和 decode 进度，拆分剩余工作与排队等待。"),
        ("10.05 - 10.11", "推进原生对照与系统性能优化",
         "优化增量备份、缺失页恢复、恢复后准入和 child 容量交接。"),
    )
    line(slide, 2.91, 2.03, 2.91, 5.40, RULE, width=2)
    for index, (date, title, detail) in enumerate(milestones):
        y = 1.80 + index * 0.84
        color = TEAL if index == 4 else BLUE
        text(slide, date, 0.78, y + 0.16, 1.96, 0.42,
             size=19, color=color, bold=True)
        marker = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(2.85), Inches(y + 0.23),
            Inches(0.12), Inches(0.12)
        )
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor.from_string(color)
        marker.line.fill.background()
        remove_shadow(marker)
        text(slide, title, 3.24, y, 9.25, 0.43,
             size=22, color=color, bold=True)
        text(slide, detail, 3.24, y + 0.43, 9.25, 0.38,
             size=18, color=GRAY)
    line(slide, 0.78, 6.08, 12.49, 6.08, RULE)
    text(slide, "目前：机制链路已贯通，稳定的端到端性能优势仍待验证。",
         0.78, 6.25, 11.71, 0.34, size=19, color=NAVY, bold=True)


def goals_slide(presentation, layout):
    slide = add_slide(
        presentation, layout, "下月目标",
        "建议用时：50 秒。\n\n"
        "下月围绕端到端收益集中推进三件事。第一，提高真正有用的预取覆盖。"
        "继续优化子任务和工具完成预测，让更多缺失的 FULL 页在返回之前恢复，"
        "而且确实被下一次执行复用。\n"
        "第二，把恢复和调度协同起来。协调可回收容量、驻留保护与执行准入，"
        "减少恢复后等待和重复加载，同时降低控制路径的额外开销。\n"
        "第三，形成公平、稳定的性能证据。采用固定配置与原生 SGLang 对照，"
        "正式实验做多轮汇总，同时观察吞吐、工作流完成时间和重算。\n"
        "最终验收关注繁忙阶段和成功工作流的吞吐提升，以及完成时间下降，"
        "而不是只看预取次数或传输量。"
    )
    text(slide, "聚焦：让预测式 KV 迁移带来可验证的端到端收益",
         0.78, 1.87, 11.73, 0.55, size=24, color=BLUE, bold=True)
    goals = (
        ("提高有效预取覆盖",
         "优化 child / tool 完成预测，让更多缺失 FULL 页在返回前恢复并被复用。"),
        ("恢复与调度协同，降低额外开销",
         "协调容量回收、驻留和准入，减少重复加载、恢复后排队与控制成本。"),
        ("建立公平、稳定的性能证据",
         "固定配置对比原生 SGLang，多轮汇总吞吐、完成时间与重算情况。"),
    )
    for index, (title, detail) in enumerate(goals, 1):
        y = 2.77 + (index - 1) * 1.10
        text(slide, f"{index:02}", 0.82, y - 0.02, 0.73, 0.60,
             size=28, color=BLUE, bold=True)
        text(slide, title, 1.86, y, 10.60, 0.48,
             size=24, color=NAVY, bold=True)
        text(slide, detail, 1.86, y + 0.53, 10.60, 0.43,
             size=20, color=GRAY)
    line(slide, 0.78, 6.08, 12.49, 6.08, RULE)
    text(slide, "验收：提高繁忙窗口与成功工作流吞吐，同时降低完成时间。",
         0.78, 6.25, 11.71, 0.34, size=18, color=TEAL, bold=True)


def main():
    presentation = Presentation(TEMPLATE)
    layout = presentation.slides[1].slide_layout
    # Keep the template's master, theme, header and logo; replace all old slides.
    for slide_id in list(presentation.slides._sldIdLst):
        presentation.part.drop_rel(slide_id.rId)
        presentation.slides._sldIdLst.remove(slide_id)
    project_slide(presentation, layout)
    progress_slide(presentation, layout)
    goals_slide(presentation, layout)
    presentation.core_properties.title = "BeliefKV 月度汇报"
    presentation.core_properties.subject = "2026 年 8 月 15 日至 10 月 11 日"
    presentation.core_properties.keywords = "BeliefKV, 动态 Agent, KV 缓存, 预测式调度"
    presentation.core_properties.comments = "三页正文；演讲者备注合计约三分钟。"
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
