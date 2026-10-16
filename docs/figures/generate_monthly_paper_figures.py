#!/usr/bin/env python3
"""Generate bilingual vector figures and native editable PowerPoint diagrams."""

from __future__ import annotations

import argparse
from collections import Counter
from functools import lru_cache
import hashlib
from html import escape
from io import BytesIO
import json
import math
from pathlib import Path
import re
import subprocess
import xml.etree.ElementTree as ET

import cairosvg
from lxml import etree
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from generate_beliefkv_paper_figures import Figure


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "docs/reports/2026-10/assets"
WIDTH = 1800
HEIGHT = 1200
PAPER_WIDTH_INCHES = 7.2
DPI = 600
FONT_ROOT = Path("/usr/share/fonts")
ARCHITECTURE_DATE = "2026-10-15"
# All explicit drawing colors are exact samples from the user's reference.
# Role, storage location and ownership use independent visual channels.
REFERENCE_PALETTE_PATH = Path(__file__).with_name("janus_reference_palette.json")
REFERENCE_PALETTE = json.loads(REFERENCE_PALETTE_PATH.read_text())
INK = "#000000"
MUTED = INK
RULE = "#a6a6a6"
PAPER = "#ffffff"
GRAY = "#a6a6a6"
GRAY_PALE = "#f2f2f2"
BACKPLATE = "#d9d9d9"
FLOW = "#156082"
BLUE = "#215f9a"
BLUE_PALE = "#a6caec"
TEAL = "#275317"
TEAL_PALE = "#84e291"
AMBER = "#35505d"
AMBER_PALE = "#f6c6ad"
PINK = "#a6caec"
YELLOW = "#f2aa84"
ACTION = TEAL
CONTROL_EDGE = "#156082"
CONTROL_PALE = "#dceaf7"
DATA_EDGE = "#275317"
DATA_PALE = "#d9f2d0"
CONTROL_HEADER = "#a6caec"
LIFECYCLE_HEADER = "#d9f2d0"
COORDINATION_HEADER = "#f6c6ad"
HBM_EDGE = INK
HBM_HEADER = "#d7aeff"
HBM_BG = "#f2f2f2"
HBM_PAGE = "#d7aeff"
HOST_EDGE = INK
HOST_HEADER = "#ffd579"
HOST_BG = "#fbe3d6"
HOST_PAGE = "#ffd579"
STORAGE_STYLES = {
    "HBM": {"edge": HBM_EDGE, "header": HBM_HEADER, "body": HBM_BG, "page": HBM_PAGE},
    "Host": {"edge": HOST_EDGE, "header": HOST_HEADER, "body": HOST_BG, "page": HOST_PAGE},
}
PEACH_PALE = GRAY_PALE
APRICOT = GRAY_PALE
LAVENDER = BLUE_PALE
MAUVE = TEAL_PALE
ORANGE_BAND = BACKPLATE
GREEN_BAND = TEAL_PALE
MATH_FAMILY = "Liberation Serif"
PALETTE = {
    "reference": "Exact pixel colors from the user-provided JANUS image; semantic mapping is BeliefKV-specific",
    "ink": INK, "gray_backplate": GRAY_PALE, "gray_panel": BACKPLATE,
    "control_header": CONTROL_EDGE, "native_header": DATA_EDGE, "blue_header": CONTROL_HEADER,
    "blue_section": BLUE_PALE, "green_section": GREEN_BAND,
    "control_plane": CONTROL_PALE, "data_plane": DATA_PALE,
    "forecast": LAVENDER, "admission": PAPER, "planner": ORANGE_BAND,
    "prepare": PAPER, "host": STORAGE_STYLES["Host"], "hbm": STORAGE_STYLES["HBM"],
    "parent": PINK, "child_1": TEAL_PALE, "child_2": AMBER_PALE,
    "transfer_engine": BACKPLATE, "transfer_band": YELLOW, "transfer_arrow": AMBER,
    "execution_arrow": FLOW, "control_action": ACTION, "feedback_arrow": BLUE,
    "lifecycle_header": LIFECYCLE_HEADER, "coordination_header": COORDINATION_HEADER,
}


def vector_commands(data: str, offset_y: float = 0):
    """Parse the M/L/C/Z paths emitted by this generator."""
    tokens = re.findall(r"[MLCZ]|[-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?", data)
    commands = []
    index = 0
    while index < len(tokens):
        command = tokens[index]
        index += 1
        count = {"M": 2, "L": 2, "C": 6, "Z": 0}[command]
        values = [float(value) for value in tokens[index:index + count]]
        if len(values) != count:
            raise ValueError("Incomplete vector path")
        points = [(values[i], values[i + 1] + offset_y) for i in range(0, count, 2)]
        commands.append((command, points))
        index += count
    return commands


ARCHITECTURE_SOURCES = (
    "beliefkv/control/causal_graph.py",
    "beliefkv/predictor/child_report_phase.py",
    "beliefkv/predictor/child_semantic_work.py",
    "beliefkv/predictor/conditional_work.py",
    "beliefkv/predictor/project_tool_history.py",
    "beliefkv/runtime/semantic_report_worker.py",
    "beliefkv/runtime/sglang_v0520_prediction.py",
    "beliefkv/runtime/sglang_v0520_runtime.py",
    "beliefkv/runtime/sglang_v0520_physical.py",
    "beliefkv/runtime/sglang_v0520_admission.py",
    "beliefkv/runtime/native_transfer_policy.py",
    "beliefkv/runtime/native_transfer_service.py",
    "patches/sglang-v0.5.20-beliefkv-staging.patch",
)


@lru_cache(maxsize=256)
def drawing_font(
    language: str, size: float, weight: int = 400,
    family: str | None = None, italic: bool = False,
):
    if family in (MATH_FAMILY, "Liberation Sans"):
        style = (
            "BoldItalic" if weight >= 600 and italic else
            "Bold" if weight >= 600 else "Italic" if italic else "Regular"
        )
        font_name = "LiberationSerif" if family == MATH_FAMILY else "LiberationSans"
        return ImageFont.truetype(
            str(FONT_ROOT / "truetype/liberation" / f"{font_name}-{style}.ttf"),
            round(size * 4),
        )
    if language == "zh":
        filename = "NotoSansCJK-Bold.ttc" if weight >= 600 else "NotoSansCJK-Regular.ttc"
        return ImageFont.truetype(
            str(FONT_ROOT / "opentype/noto" / filename), round(size * 4), index=2,
        )
    filename = (
        "DejaVuSans-BoldOblique.ttf" if weight >= 600 and italic else
        "DejaVuSans-Bold.ttf" if weight >= 600 else
        "DejaVuSans-Oblique.ttf" if italic else "DejaVuSans.ttf"
    )
    return ImageFont.truetype(
        str(FONT_ROOT / "truetype/dejavu" / filename), round(size * 4),
    )


class PaperFigure(Figure):
    def __init__(
        self, name: str, title: str, height: int, language: str, *, width=WIDTH,
    ):
        super().__init__(name + "_" + language + ".svg", title, width, height)
        self.language = language
        self.family = "Noto Sans CJK SC" if language == "zh" else "DejaVu Sans"
        self.parts = [
            '<rect width="100%" height="100%" fill="#ffffff"/>',
            "<defs>",
            *(self.flow_marker(name, color, kind, size)
              for name, color, kind in (
                  ("execution", FLOW, "triangle"), ("feedback", BLUE, "open"),
                  ("action", ACTION, "notched"), ("transfer", AMBER, "triangle"),
                  ("inactive", GRAY, "open"),
              ) for size in (12, 18, 24, 32)),
            "</defs>",
            "<style>text{letter-spacing:0}"
            "path,line{stroke-linejoin:round;stroke-linecap:round}</style>",
        ]
        self.text_bounds = []
        self.math_objects = 0
        self.panel_styles = []
        self.flow_types = Counter()
        self.flow_geometry = Counter()
        self.transfer_entities = []
        self.transfer_actions = 0
        self.execution_steps = []
        self.memory_tiers = []
        self.storage_pages = []

    @staticmethod
    def flow_marker(name, color, kind, size):
        glyph = (
            f'<path d="M2 1L16 8L2 15" fill="none" stroke="{color}" '
            'stroke-width="2.5" stroke-linejoin="miter"/>'
            if kind == "open" else
            f'<path d="M0 0L16 8L0 16Z" fill="{color}"/>'
            if kind == "triangle" else
            f'<path d="M0 0L16 8L0 16L4 8Z" fill="{color}"/>'
        )
        return (
            f'<marker id="{name}-{size}" markerUnits="userSpaceOnUse" '
            f'viewBox="0 0 16 16" refX="16" refY="8" '
            f'markerWidth="{size}" markerHeight="{size}" orient="auto">'
            f'{glyph}</marker>'
        )

    def tr(self, english: str, chinese: str) -> str:
        return chinese if self.language == "zh" else english

    def line(self, x1, y1, x2, y2, color=INK, width=1.7, dash=None, arrow=None):
        super().line(x1, y1, x2, y2, color, width, dash, arrow)

    def path(self, points, color=INK, width=1.7, dash=None, arrow=None):
        super().path(points, color, width, dash, arrow)

    def rect(self, x, y, w, h, fill=PAPER, stroke=RULE, width=1.3, radius=2, dash=None):
        super().rect(x, y, w, h, fill, stroke, width, radius, dash)

    def circle(self, x, y, r, fill=PAPER, stroke=INK, width=1.5):
        super().circle(x, y, r, fill, stroke, width)

    def diamond(self, cx, cy, w, h, fill=AMBER_PALE, stroke=AMBER):
        super().diamond(cx, cy, w, h, fill, stroke)

    def measure(
        self, value: str, size: float, weight: int = 400,
        family: str | None = None, italic: bool = False,
    ):
        font = drawing_font(self.language, size, weight, family, italic)
        _, top, _, bottom = font.getbbox(value, anchor="ls")
        return font.getlength(value) / 4, (bottom - top) / 4, -top / 4

    def text(
        self, x, y, value, size=24, color=INK, weight=400,
        anchor="start", family=None, italic=False,
    ):
        family = family or (
            "Liberation Sans" if self.language == "en" and italic else self.family
        )
        color = PAPER if color == PAPER else INK
        width, height, baseline = self.measure(value, size, weight, family, italic)
        left = x - (width / 2 if anchor == "middle" else width if anchor == "end" else 0)
        self.text_bounds.append((value, left, y - baseline, width, height))
        self.parts.append(
            f'<text x="{x}" y="{y}" xml:space="preserve" font-size="{size}" fill="{color}" '
            f'font-weight="{weight}" text-anchor="{anchor}" '
            f'font-family="{escape(family)}" font-style="{"italic" if italic else "normal"}">'
            f"{escape(value)}</text>"
        )
        if family == MATH_FAMILY:
            self.math_objects += 1

    def panel(
        self, x, y, w, h, title, *, color=BLUE_PALE, fill=GRAY_PALE,
        size=26, band=38, border=RULE, title_color=INK, role=None,
    ):
        self.rect(x, y, w, h, fill, border, 2, 8)
        if role:
            self.parts[-1] = self.parts[-1].replace(
                "<rect ", f'<rect data-container-role="{role}" ', 1,
            )
        self.rect(x + 1, y + 1, w - 2, band - 1, color, "none", 0, 7)
        self.rect(x + 1, y + band / 2, w - 2, band / 2, color, "none", 0, 0)
        self.line(x + 1, y + band, x + w - 1, y + band, border, 1.2)
        inset = 14
        title_width = self.measure(title, size, 700)[0]
        if title_width > w - inset - 14:
            size *= (w - inset - 14) / title_width
        self.text(x + inset, y + band * .72 + (4 if self.language == "zh" else 0),
                  title, size, title_color, 700)
        self.panel_styles.append({
            "title": title, "header": color, "body": fill, "border": border,
            "position": [x, y], "size": [w, h], "radius": 8, "role": role,
        })

    def heading(self, english, chinese, subtitle_en, subtitle_zh):
        self.text(40, 51, self.tr(english, chinese), 37, INK, 700)
        self.text(40, 88, self.tr(subtitle_en, subtitle_zh), 24, MUTED)
        self.line(40, 108, self.width - 40, 108, INK, 1.6)

    def mechanism_title(self, design, english, chinese):
        self.text(32, 43, design, 31, INK, 700)
        self.text(108, 43, self.tr(english, chinese), 31, INK, 700)

    def memory_pool(self, x, y, w, h, title, tier):
        style = STORAGE_STYLES[tier]
        self.rect(x, y, w, h, style["body"], style["edge"], 2.2, 8)
        self.parts[-1] = self.parts[-1].replace(
            "<rect ", f'<rect data-memory-tier="{tier}" ', 1,
        )
        self.rect(x + 1, y + 1, w - 2, 37, style["header"], "none", 0, 7)
        self.parts[-1] = self.parts[-1].replace(
            "<rect ", f'<rect data-tier-header="{tier}" ', 1,
        )
        self.rect(x + 1, y + 20, w - 2, 18, style["header"], "none", 0, 0)
        self.line(x + 1, y + 38, x + w - 1, y + 38, style["edge"], 1.2)
        self.text(x + 16, y + 28, title, 27, INK, 700)
        self.memory_tiers.append({
            "tier": tier, "position": [x, y], "size": [w, h],
            "header": style["header"], "body": style["body"], "border": style["edge"],
            "role": "peer_storage_tier",
        })

    def storage_band(self, x, y, w, h, tier):
        style = STORAGE_STYLES[tier]
        self.rect(x, y, w, h, style["body"], style["edge"], 1.8, 8)
        self.parts[-1] = self.parts[-1].replace(
            "<rect ", f'<rect data-storage-band="{tier}" ', 1,
        )

    def storage_page(
        self, x, y, label="", *, tier, owner=PINK, missing=False,
        w=54, h=47, size=21,
    ):
        style = STORAGE_STYLES[tier]
        self.rect(x, y, w, h, PAPER if missing else style["page"],
                  style["edge"], 1.6, 2, "6 4" if missing else None)
        self.parts[-1] = self.parts[-1].replace(
            "<rect ", f'<rect data-page-tier="{tier}" data-page-owner="{owner}" '
            f'data-page-missing="{str(missing).lower()}" ', 1,
        )
        stripe_height = min(5, h * .15)
        self.rect(x + 2, y + 2, w - 4, stripe_height, owner, "none", 0, 1)
        if label:
            self.centered(x + w / 2, y + h / 2 + size * .4,
                          label, size=size, bold=True)
        self.storage_pages.append({
            "tier": tier, "owner": owner, "missing": missing, "position": [x, y],
        })

    def storage_state(self, x, y, *, tier, owner=PINK, scale=.6, exists=True):
        self.snapshot(x, y, fill=STORAGE_STYLES[tier]["page"] if exists else PAPER,
                      scale=scale)
        if exists:
            self.rect(x + 14 * scale, y + 7 * scale, 51 * scale, 5 * scale,
                      owner, "none", 0, 1)

    def step(self, x, y, number, meaning):
        self.circle(x, y, 19, INK, PAPER, 3)
        self.parts[-1] = self.parts[-1].replace(
            "<circle ", f'<circle data-execution-step="{number}" ', 1,
        )
        self.centered(x, y + 7.5, str(number), size=23, color=PAPER, bold=True)
        self.execution_steps.append({
            "number": number, "position": [x, y], "meaning": meaning,
        })

    def _route(self, commands, *, color, dashed, width, geometry):
        if color == INK:
            color = FLOW
        role = {
            FLOW: "execution", BLUE: "feedback", ACTION: "action",
            AMBER: "transfer", GRAY: "inactive",
        }[color]
        kind = (
            "open" if role in ("feedback", "inactive") else
            "notched" if role == "action" else "triangle"
        )
        minimum_width = 4.0 if kind == "open" else 8.5 if role == "transfer" else 4.8
        width = max(width, minimum_width)
        points = [point for _, part in commands for point in part]
        last_length = math.dist(points[-2], points[-1])
        size = (
            12 if last_length < 28 else
            24 if role == "transfer" and last_length < 50 else
            32 if role == "transfer" else
            24 if role == "action" and last_length > 70 else 18
        )
        data = " ".join(command + " ".join(f"{x:.4f},{y:.4f}" for x, y in part)
                        for command, part in commands)
        self.parts.append(
            f'<path d="{data}" fill="none" stroke="{PAPER}" stroke-width="{width + 2}"/>'
        )
        dash = ' stroke-dasharray="12 8"' if dashed else ""
        self.parts.append(
            f'<path d="{data}" fill="none" stroke="{color}" stroke-width="{width}"{dash} '
            f'marker-end="url(#{role}-{size})" data-arrow-kind="{kind}" '
            f'data-arrow-size="{size}" data-flow-role="{role}" '
            f'data-flow-geometry="{geometry}"/>'
        )
        self.flow_types[role] += 1
        self.flow_geometry[geometry] += 1

    def arrow(self, points, *, color=FLOW, dashed=False, width=4.8):
        commands = [("M", [points[0]])]
        for index, point in enumerate(points[1:-1], 1):
            previous, following = points[index - 1], points[index + 1]
            before, after = math.dist(previous, point), math.dist(point, following)
            if min(before, after) == 0:
                continue
            radius = min(18, before / 3, after / 3)
            incoming = tuple((point[i] - previous[i]) / before for i in (0, 1))
            outgoing = tuple((following[i] - point[i]) / after for i in (0, 1))
            start = tuple(point[i] - incoming[i] * radius for i in (0, 1))
            end = tuple(point[i] + outgoing[i] * radius for i in (0, 1))
            control1 = tuple(start[i] + incoming[i] * radius * .5523 for i in (0, 1))
            control2 = tuple(end[i] - outgoing[i] * radius * .5523 for i in (0, 1))
            commands.extend((("L", [start]), ("C", [control1, control2, end])))
        commands.append(("L", [points[-1]]))
        self._route(commands, color=color, dashed=dashed, width=width,
                    geometry="rounded" if len(points) > 2 else "straight")

    def curve_arrow(self, start, control1, control2, end, *,
                    color=FLOW, dashed=False, width=5.5):
        self._route([("M", [start]), ("C", [control1, control2, end])],
                    color=color, dashed=dashed, width=width, geometry="bezier")

    def transfer_curve(self, start, control1, control2, end, *, width=17):
        def unit(delta):
            length = math.hypot(*delta)
            return tuple(value / length for value in delta)

        forward0 = unit(tuple(control1[i] - start[i] for i in (0, 1)))
        forward1 = unit(tuple(end[i] - control2[i] for i in (0, 1)))
        normal0, normal1 = (-forward0[1], forward0[0]), (-forward1[1], forward1[0])
        shoulder = tuple(end[i] - forward1[i] * 28 for i in (0, 1))
        bend = tuple(shoulder[i] - forward1[i] * 38 for i in (0, 1))

        def offset(point, normal, distance):
            return tuple(point[i] + normal[i] * distance for i in (0, 1))

        half = width / 2
        commands = [
            ("M", [offset(start, normal0, half)]),
            ("C", [offset(control1, normal0, half), offset(bend, normal1, half),
                   offset(shoulder, normal1, half)]),
            ("L", [offset(shoulder, normal1, width)]),
            ("L", [end]),
            ("L", [offset(shoulder, normal1, -width)]),
            ("L", [offset(shoulder, normal1, -half)]),
            ("C", [offset(bend, normal1, -half), offset(control1, normal0, -half),
                   offset(start, normal0, -half)]),
            ("Z", []),
        ]
        data = " ".join(command + " ".join(f"{x:.4f},{y:.4f}" for x, y in points)
                        for command, points in commands)
        self.parts.append(
            f'<path d="{data}" fill="{YELLOW}" stroke="{ACTION}" stroke-width="1.5" '
            'data-semantic-role="transfer-action" data-flow-geometry="curved_ribbon"/>'
        )
        self.flow_types["transfer"] += 1
        self.flow_geometry["curved_ribbon"] += 1

    def operation_label(self, x, y, label, *, size=24):
        self.centered(x, y, label, size=size, bold=True)
        self.parts[-1] = self.parts[-1].replace(
            "<text ", '<text data-semantic-role="transfer-action" ', 1,
        )
        self.transfer_actions += 1

    def transfer_band(self, x1, y, x2, *, h=25, label=None, allocation=False):
        direction = 1 if x2 > x1 else -1
        shoulder = x2 - direction * min(19, abs(x2 - x1) * .35)
        self.polygon((
            (x1, y - h * .34), (shoulder, y - h * .34),
            (shoulder, y - h * .65), (x2, y),
            (shoulder, y + h * .65), (shoulder, y + h * .34),
            (x1, y + h * .34),
        ), BACKPLATE if allocation else YELLOW, FLOW if allocation else ACTION, 1.4)
        self.flow_types["capacity_allocation" if allocation else "transfer_band"] += 1
        self.flow_geometry["wide_block"] += 1
        if label:
            if allocation:
                self.centered((x1 + x2) / 2 - 4, y + 7, label, size=min(23, h * .75))
            else:
                self.operation_label((x1 + x2) / 2, y + 7, label, size=min(25, h * .8))

    def activity_bar(self, x, y, w, h, *, label=None):
        head = min(17, w / 4)
        self.polygon(((x, y), (x + w - head, y), (x + w, y + h / 2),
                      (x + w - head, y + h), (x, y + h)), YELLOW, ACTION, 1.2)
        self.parts[-1] = self.parts[-1].replace(
            "<polygon ", '<polygon data-semantic-role="transfer-duration" ', 1,
        )
        if label:
            self.operation_label(x + w / 2, y + h / 2 + 8, label, size=min(25, h * .73))

    def centered(self, x, y, value, *, size=23, color=INK, bold=False):
        self.text(x, y, value, size, color, 700 if bold else 400, "middle")

    def page(
        self, x, y, label="", *, fill=PINK, stroke=INK, dashed=False,
        w=54, h=47, size=21,
    ):
        self.rect(x, y, w, h, fill, stroke, 1.8, 1, "6 4" if dashed else None)
        if label:
            self.centered(x + w / 2, y + h / 2 + size * .34,
                          label, size=size, bold=True)

    def token(self, x, y, label, fill=PINK, stroke=INK, r=27, size=25):
        self.circle(x, y, r, fill, stroke, 2)
        self.centered(x, y + size * .34, label, size=size, bold=True)

    def badge(self, x, y, label, *, fill=PAPER, r=16):
        self.circle(x, y, r, fill, INK, 1.8)
        self.centered(x, y + 6, str(label), size=21, bold=True)

    def math_width(self, tokens, size=30):
        total = 0
        for token in tokens:
            if isinstance(token, str):
                total += self.measure(token, size, family=MATH_FAMILY)[0]
            else:
                base, sub, sup = (*token, None, None)[:3]
                total += self.measure(base, size, family=MATH_FAMILY, italic=True)[0]
                if sub or sup:
                    total += max(
                        self.measure(value or "", size * .65, family=MATH_FAMILY)[0]
                        for value in (sub, sup)
                    ) + 2
        return total

    def math(self, x, y, tokens, *, size=30, color=INK, anchor="start"):
        width = self.math_width(tokens, size)
        x -= width / 2 if anchor == "middle" else width if anchor == "end" else 0
        for token in tokens:
            if isinstance(token, str):
                self.text(x, y, token, size, color, family=MATH_FAMILY)
                x += self.measure(token, size, family=MATH_FAMILY)[0]
                continue
            base, sub, sup = (*token, None, None)[:3]
            self.text(x, y, base, size, color, family=MATH_FAMILY, italic=True)
            x += self.measure(base, size, family=MATH_FAMILY, italic=True)[0]
            if sub:
                self.text(x + 1, y + size * .25, sub, size * .65, color,
                          family=MATH_FAMILY)
            if sup:
                self.text(x + 1, y - size * .52, sup, size * .65, color,
                          family=MATH_FAMILY)
            if sub or sup:
                x += max(
                    self.measure(value or "", size * .65, family=MATH_FAMILY)[0]
                    for value in (sub, sup)
                ) + 2
        return width

    def fraction(self, x, y, numerator, denominator, *, size=30, color=INK):
        width = max(self.math_width(tokens, size) for tokens in (numerator, denominator)) + 20
        self.math(x + width / 2, y - 9, numerator, size=size, color=color, anchor="middle")
        self.line(x, y, x + width, y, color, 1.8)
        self.math(x + width / 2, y + size + 3, denominator,
                  size=size, color=color, anchor="middle")
        return width

    def document(self, x, y, w=130, h=112, *, color=BLUE_PALE, labels=()):
        self.rect(x, y, w, h, PAPER, INK, 1.9, 2)
        self.polygon(((x + w - 24, y), (x + w, y + 24), (x + w - 24, y + 24)),
                     PAPER, INK)
        for index, fraction in enumerate((.72, .84, .65, .77)):
            yy = y + 33 + index * 19
            if yy > y + h - 10:
                break
            if index < len(labels):
                self.text(x + 10, yy + 1, labels[index], min(19, w * .135), INK)
            else:
                self.rect(x + 12, yy - 9, (w - 25) * fraction, 10,
                          color if index in (1, 3) else BACKPLATE, "none", 0, 0)

    def polygon(self, points, fill=PAPER, stroke=INK, width=1.8):
        self.parts.append(
            '<polygon points="' + " ".join(f"{x},{y}" for x, y in points)
            + f'" fill="{fill}" stroke="{stroke}" stroke-width="{width}"/>'
        )

    def bank(self, x, y, w, h, colors, *, label=None, missing=()):
        self.rect(x, y, w, h, PAPER, INK, 2, 3)
        self.line(x + 8, y + 17, x + w - 8, y + 17, INK, 1)
        gap = 6
        page_w = (w - 20 - (len(colors) - 1) * gap) / len(colors)
        for index, color in enumerate(colors):
            self.page(x + 10 + index * (page_w + gap), y + 29,
                      fill=color if index not in missing else PAPER,
                      dashed=index in missing, w=page_w, h=h - 43)
        for xx in range(round(x + 13), round(x + w - 5), 17):
            self.rect(xx, y + h, 7, 8, YELLOW, INK, 1, 0)
        if label:
            self.centered(x + w / 2, y - 15, label, size=27, bold=True)

    def snapshot(self, x, y, label="", *, fill=PINK, scale=1, faded=False):
        self.polygon(tuple((x + xx * scale, y + yy * scale) for xx, yy in (
            (12, 0), (67, 0), (79, 29.5), (67, 59), (12, 59), (0, 29.5),
        )), PAPER if faded else fill, GRAY if faded else INK, 1.5)
        self.centered(x + 39.5 * scale, y + 40 * scale,
                      "S" if fill != PAPER or faded else "∅",
                      size=33 * scale, color=GRAY if faded else INK, bold=True)
        if label:
            self.math(x + 40 * scale, y + 87 * scale, [label],
                      size=27 * scale, color=GRAY if faded else INK, anchor="middle")

    def table(self, x, y, widths, rows, *, row_h=30, size=20, header=True):
        for row_index, values in enumerate(rows):
            xx = x
            for width, value in zip(widths, values):
                self.rect(xx, y + row_index * row_h, width, row_h,
                          BLUE_PALE if header and row_index == 0 else PAPER,
                          INK, 1, 0)
                self.text(xx + 8, y + row_index * row_h + row_h * .71,
                          value, size, INK, 700 if header and row_index == 0 else 400)
                xx += width

    def strip(self, x, y, fills, *, labels=None, w=40, h=30, gap=5, missing=()):
        for index, fill in enumerate(fills):
            self.page(x + index * (w + gap), y,
                      labels[index] if labels else "",
                      fill=PAPER if index in missing else fill,
                      stroke=GRAY if index in missing else INK,
                      dashed=index in missing, w=w, h=h, size=min(21, h * .65))

    def bracket(self, x1, x2, y, *, color=INK, down=False):
        direction = -1 if down else 1
        self.path(((x1, y + 7 * direction), (x1, y),
                   (x2, y), (x2, y + 7 * direction)), color, 1.6)

    def gpu(self, x, y, w=130, h=114, *, label="GPU", running=None):
        for xx in range(round(x + 14), round(x + w - 5), 20):
            self.line(xx, y - 10, xx, y, INK, 2)
            self.line(xx, y + h, xx, y + h + 10, INK, 2)
        for yy in range(round(y + 15), round(y + h - 5), 20):
            self.line(x - 10, yy, x, yy, INK, 2)
            self.line(x + w, yy, x + w + 10, yy, INK, 2)
        self.rect(x, y, w, h, BLUE_PALE, INK, 2, 2)
        self.centered(x + w / 2, y + 29, label, size=24, bold=True)
        rows, columns = 2, 3
        for row in range(rows):
            for column in range(columns):
                index = row * columns + column
                fill = running[index] if running and index < len(running) else PAPER
                self.rect(x + 13 + column * (w - 22) / columns,
                          y + 42 + row * (h - 53) / rows,
                          (w - 35) / columns, (h - 66) / rows,
                          fill, INK, 1.1, 1)

    def transfer_engine(self, x, y, w=142, h=85, *, label="DMA"):
        if label in ("D2H", "H2D"):
            raise ValueError("Transfer directions label edges, not component nodes")
        self.transfer_entities.append(label)
        inset = min(12, h * .18)
        self.polygon(((x + inset, y), (x + w - inset, y), (x + w, y + inset),
                      (x + w, y + h - inset), (x + w - inset, y + h),
                      (x + inset, y + h), (x, y + h - inset), (x, y + inset)),
                     BACKPLATE, INK, 1.8)
        self.centered(x + w / 2, y + min(29, h * .45) + (4 if self.language == "zh" else 0),
                      label, size=26, bold=True)
        for index in range(3):
            yy = y + h * .55 + index * 7
            self.line(x - 8, yy, x + 7, yy, FLOW, 2)
            self.line(x + w - 7, yy, x + w + 8, yy, FLOW, 2)

    def encoder(self, x, y, w=166, h=126):
        for yy in (y + 29, y + h - 29):
            self.line(x - 10, yy, x, yy, FLOW, 2)
            self.line(x + w, yy, x + w + 10, yy, FLOW, 2)
        self.rect(x, y, w, h, LAVENDER, INK, 1.8, 5)
        self.text(x + 18, y + 36, "Text", 24, INK, 700)
        self.text(x + 18, y + 67, "encoder", 24, INK, 700)
        self.lock(x + w - 35, y + 15, scale=.7)
        self.math(x + w / 2, y + h - 11, [("f", "θ")], size=25, anchor="middle")
        self.centered(x + w / 2, y + h + 32, self.tr("Frozen weights", "冻结权重"),
                      size=21)

    def probability_bars(self, x, y, *, width=178, row_h=36, size=19):
        for index, (label, fill, fraction) in enumerate((
            (self.tr("work", "继续"), BLUE_PALE, .1),
            (self.tr("notice", "通知"), AMBER_PALE, .2),
            (self.tr("final", "报告"), PINK, .7),
        )):
            yy = y + index * row_h
            self.text(x, yy + 17, label, size)
            self.rect(x + 73, yy + 2, width - 73, 18, PAPER, RULE, .8, 0)
            self.rect(x + 73, yy + 2, (width - 73) * fraction, 18, fill, INK, .9, 0)

    def ticket(self, x, y, label, fill, *, w=180, h=52, resident=False, protected=False):
        self.rect(x, y, w, h, PAPER, INK, 1.8, 2)
        self.rect(x + 1, y + 1, 10, h - 2, fill, "none", 0, 0)
        size = 19 if h < 40 else 23
        self.centered(x + 48, y + h / 2 + size * .34, label, size=size, bold=True)
        for index in range(3):
            page_h = min(22, h - 8)
            self.storage_page(x + w - 75 + index * 20, y + (h - page_h) / 2,
                              tier="HBM", owner=fill, missing=not resident,
                              w=15, h=page_h)
        if protected:
            self.lock(x + w - 20, y - 18, color=INK, scale=.6)

    def clock(self, x, y, radius=25, *, color=INK):
        self.circle(x, y, radius, PAPER, color, 1.8)
        self.path(((x, y - radius * .65), (x, y), (x + radius * .55, y + radius * .25)),
                  color, 2)

    def terminal(self, x, y, w=105, h=65):
        self.rect(x, y, w, h, PAPER, INK, 1.8, 3)
        self.line(x + 1, y + h * .21, x + w - 1, y + h * .21, INK, 1)
        self.path(((x + 16, y + h * .43), (x + 27, y + h * .61),
                   (x + 16, y + h * .79)), INK, 2)
        self.line(x + 38, y + h * .78, x + w - 14, y + h * .78, INK, 2)

    def junction(self, x, y, *, kind="and", radius=18):
        self.circle(x, y, radius, PAPER, INK, 1.6)
        if kind == "concat":
            for offset in (-4, 4):
                self.line(x + offset, y - 10, x + offset, y + 10, INK, 1.8)
        else:
            size = min(23, radius * 1.15)
            self.centered(x, y + size * .33, "&", size=size, bold=True)

    def lock(self, x, y, *, color=INK, scale=1):
        self.path(((x + 5 * scale, y + 13 * scale),
                   (x + 5 * scale, y + 2 * scale),
                   (x + 19 * scale, y + 2 * scale),
                   (x + 19 * scale, y + 13 * scale)), color, 2)
        self.rect(x, y + 13 * scale, 24 * scale, 20 * scale, PAPER, color, 1.8, 2)
        self.circle(x + 12 * scale, y + 21 * scale, 2 * scale, color, color, 1)

    def check(self, x, y, *, color=TEAL, size=20):
        self.path(((x, y + size * .5), (x + size * .35, y + size),
                   (x + size, y)), color, 2.8)

    def cross(self, x, y, *, color=GRAY, size=20):
        self.line(x, y, x + size, y + size, color, 2.8)
        self.line(x, y + size, x + size, y, color, 2.8)

    def interval(self, x1, x2, y, *, color=TEAL, center=None):
        self.line(x1, y, x2, y, color, 2.5)
        self.line(x1, y - 9, x1, y + 9, color, 2)
        self.line(x2, y - 9, x2, y + 9, color, 2)
        self.circle(center or (x1 + x2) / 2, y, 5, color, color, 1)

    def design_footer(self, active=None):
        self.line(40, 1118, 1760, 1118, RULE, 1.2)
        for number, x, label, color in (
            (1, 64, self.tr("Demand forecast", "执行需求预测"), BLUE_PALE),
            (3, 775, self.tr("Prefetch + admission", "预取与准入"), BACKPLATE),
            (2, 1434, self.tr("KV lifecycle", "KV 生命周期"), TEAL_PALE),
        ):
            self.circle(x, 1155, 21, INK if active in (None, number) else PAPER, INK, 1.8)
            self.centered(x, 1162, f"D{number}", size=21, bold=True,
                          color=PAPER if active in (None, number) else INK)
            self.text(x + 32, 1163, label, 22, INK,
                      700 if active in (None, number) else 400)
        self.arrow(((342, 1155), (745, 1155)), color=BLUE, dashed=True)
        self.centered(540, 1140, self.tr("agent + timing", "候选 + 时机"), size=17)
        self.arrow(((1404, 1155), (1115, 1155)), color=ACTION)
        self.centered(1260, 1140, self.tr("checkpoint / pages / capacity",
                                       "checkpoint / 缺页 / 容量"), size=16)
        self.curve_arrow((1115, 1181), (1220, 1193), (1375, 1193), (1412, 1174),
                         color=ACTION)
        self.centered(1260, 1180, self.tr("physical actions", "物理动作"), size=16)
        self.arrow(((775, 1179), (775, 1194), (64, 1194), (64, 1179)),
                   color=BLUE, dashed=True)
        self.centered(414, 1185, self.tr("service / progress / completion",
                                       "服务 / 进度 / 完成事件"), size=16)

    def save_to(self, directory: Path):
        for value, x, y, width, height in self.text_bounds:
            if x < 0 or y < 0 or x + width > self.width or y + height > self.height:
                raise ValueError(f"{self.filename}: text outside canvas: {value!r}")
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / self.filename
        physical_height = PAPER_WIDTH_INCHES * self.height / self.width
        svg = (
            '<?xml version="1.0" encoding="UTF-8"?>\n'
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{PAPER_WIDTH_INCHES}in" '
            f'height="{physical_height:.6f}in" viewBox="0 0 {self.width} {self.height}" '
            'role="img" aria-labelledby="title">\n'
            f'<title id="title">{escape(self.title)}</title>\n'
            + "\n".join(self.parts) + "\n</svg>\n"
        )
        path.write_text(svg, encoding="utf-8")
        cairosvg.svg2pdf(
            bytestring=svg.encode("utf-8"), write_to=str(path.with_suffix(".pdf")),
            dpi=72, output_width=PAPER_WIDTH_INCHES * 72,
            output_height=physical_height * 72,
        )
        png_width = round(PAPER_WIDTH_INCHES * DPI)
        png_height = round(physical_height * DPI)
        rendered = cairosvg.svg2png(
            bytestring=svg.encode("utf-8"), output_width=png_width,
            output_height=png_height,
        )
        with Image.open(BytesIO(rendered)) as image:
            image.convert("RGB").save(path.with_suffix(".png"), dpi=(DPI, DPI))
        return {
            "stem": path.stem, "language": self.language, "title": self.title,
            "canvas": [self.width, self.height], "png_pixels": [png_width, png_height],
            "paper_width_inches": PAPER_WIDTH_INCHES, "dpi": DPI,
            "text_objects": len(self.text_bounds),
            "editable_math_text_objects": self.math_objects,
            "panel_styles": self.panel_styles,
            "flow_types": dict(self.flow_types),
            "flow_geometry": dict(self.flow_geometry),
            "transfer_entities": self.transfer_entities,
            "transfer_action_labels": self.transfer_actions,
            "execution_steps": self.execution_steps,
            "memory_tiers": self.memory_tiers,
            "storage_pages": self.storage_pages,
            "layout": "overview" if self.filename.startswith("01_") else "compact_algorithm",
        }


def system_overview(language: str) -> PaperFigure:
    f = PaperFigure("01_system_overview", "BeliefKV Execution and KV Coordination", HEIGHT, language)
    f.heading(
        "BeliefKV: Workflow-aware Execution and KV Management",
        "BeliefKV：工作流感知的执行与 KV 协同管理",
        "Forecast demand → manage KV → admit ready work",
        "预测执行需求 → 管理 KV 生命周期 → 协同准入",
    )
    f.panel(32, 124, 1736, 152, f.tr("Dynamic agent workflow", "动态 Agent 工作流"),
            color=BACKPLATE, fill=GRAY_PALE, border=GRAY)
    f.token(83, 220, "P", PINK, r=25, size=24)
    f.token(237, 190, "C1", TEAL_PALE, r=23, size=22)
    f.token(237, 246, "C2", AMBER_PALE, r=23, size=22)
    f.terminal(327, 172, 68, 40)
    f.diamond(487, 219, 107, 66, PAPER, INK)
    f.centered(487, 227, "JOIN", size=24, bold=True)
    f.token(620, 219, "P", PINK, r=25, size=24)
    for points in (
        ((111, 211), (210, 193)), ((111, 232), (210, 243)),
        ((263, 190), (324, 190)),
        ((398, 191), (417, 191), (433, 204)),
        ((263, 246), (414, 246), (433, 233)), ((544, 219), (592, 219)),
    ):
        f.arrow(points)
    f.text(127, 184, "SPAWN", 21)
    f.centered(437, 177, "RETURN", size=18)
    f._route([
        ("M", [(648, 224)]),
        ("C", [(728, 215), (742, 264), (663, 262)]),
        ("C", [(631, 262), (620, 269), (620, 247)]),
    ], color=FLOW, dashed=True, width=4.8, geometry="loop")
    f.centered(709, 185, f.tr("next round", "下一轮"), size=21)
    f.line(780, 171, 780, 260, RULE, 1.2)
    f.arrow(((815, 211), (1112, 211)), color=BLUE, dashed=True)
    for xx in (847, 956, 1065):
        f.circle(xx, 211, 6, BLUE, BLUE, 1)
    f.centered(957, 256, f.tr("events", "事件"), size=23)
    f.document(1190, 171, 82, 76, color=AMBER_PALE)
    f.centered(1231, 268, f.tr("text", "正文"), size=22)
    f.polygon(((1363, 181), (1485, 181), (1485, 223), (1420, 223),
               (1404, 237), (1404, 223), (1363, 223)), YELLOW, INK)
    f.centered(1424, 209, "NOTICE", size=22, bold=True)
    f.path(((1584, 236), (1606, 236), (1606, 219), (1631, 219),
            (1631, 201), (1657, 201), (1657, 182), (1689, 182)), BLUE, 2.5)
    f.centered(1637, 268, f.tr("progress", "进度"), size=22)

    f.panel(32, 300, 912, 804, f.tr("BeliefKV control plane", "BeliefKV 控制面"),
            color=CONTROL_EDGE, fill=CONTROL_PALE, border=CONTROL_EDGE,
            title_color=PAPER, size=28, role="control_plane")
    f.panel(968, 300, 800, 804, f.tr("Native SGLang data plane", "原生 SGLang 数据面"),
            color=DATA_EDGE, fill=DATA_PALE, border=DATA_EDGE,
            title_color=PAPER, size=28, role="data_plane")
    f.arrow(((1230, 279), (1230, 287), (453, 287), (453, 298)), color=BLUE, dashed=True)
    f.panel(48, 350, 880, 292, f.tr("D1  Event + content forecasts",
                                                  "D1  事件与内容驱动的需求预测"),
            color=CONTROL_HEADER, fill=PAPER, border=CONTROL_EDGE, size=26)
    f.token(91, 478, "P", PINK, r=24, size=23)
    f.token(202, 427, "C1", TEAL_PALE, r=23, size=22)
    f.token(202, 537, "C2", AMBER_PALE, r=23, size=22)
    f.check(234, 403, color=INK, size=18)
    f.diamond(323, 478, 83, 60, PAPER, INK)
    f.centered(323, 487, "ALL", size=22, bold=True)
    f.arrow(((118, 464), (176, 439)))
    f.arrow(((118, 493), (176, 524)))
    f.arrow(((230, 427), (253, 427), (282, 463)), color=GRAY)
    f.arrow(((230, 537), (253, 537), (282, 493)), color=BLUE, dashed=True)
    f.text(69, 609, "pending {C2}", 24, INK, 700)
    f.document(427, 422, 73, 76, color=AMBER_PALE)
    f.polygon(((427, 525), (517, 525), (517, 557), (471, 557),
               (458, 569), (458, 557), (427, 557)), AMBER_PALE, INK)
    f.centered(472, 549, "NOTICE", size=17, bold=True)
    f.arrow(((230, 537), (390, 537), (390, 459), (423, 459)))
    f.arrow(((504, 461), (567, 461)))
    f.arrow(((520, 541), (545, 541), (545, 461), (567, 461)))
    f.text(572, 419, f.tr("phase / work", "阶段 / 工作"), 22, INK, 700)
    f.rect(577, 438, 99, 17, PAPER, RULE, 1, 0)
    f.rect(577, 438, 70, 17, BLUE_PALE, INK, 1, 0)
    f.interval(577, 676, 483, color=INK, center=637)
    f.arrow(((685, 461), (719, 461)), color=ACTION)
    f.clock(744, 461, 19)
    f.arrow(((771, 461), (829, 461)), color=BLUE, dashed=True)
    f.token(861, 461, "P", PINK, r=23, size=22)
    f.centered(815, 507, "RETURN", size=21)
    f.arrow(((368, 478), (386, 478), (386, 395), (861, 395), (861, 434)),
            color=BLUE, dashed=True)
    f.terminal(427, 582, 67, 43)
    f.clock(534, 605, 17)
    f.path(((578, 620), (600, 620), (600, 606), (628, 606),
            (628, 591), (675, 591)), BLUE, 2.8)
    f.arrow(((557, 605), (572, 605)))
    f.arrow(((685, 605), (719, 605)), color=ACTION)
    f.interval(727, 804, 605, color=INK, center=770)
    f.arrow(((812, 605), (832, 605)), color=BLUE, dashed=True)
    f.token(861, 605, "A", TEAL_PALE, r=23, size=22)
    f.centered(769, 575, "TOOL_END", size=19)

    f.panel(48, 670, 430, 434, f.tr("D2  KV lifecycle", "D2  KV 生命周期"),
            color=LIFECYCLE_HEADER, fill=PAPER, border=CONTROL_EDGE, size=26)
    f.text(64, 749, "WAIT_JOIN / WAIT_TOOL", 21, INK, 700)
    f.token(87, 790, "P", PINK, r=20, size=20)
    f.storage_band(136, 764, 108, 88, "HBM")
    f.storage_band(306, 764, 108, 88, "Host")
    for index in range(3):
        f.storage_page(144 + index * 33, 774, tier="HBM", w=28, h=31)
        f.storage_page(314 + index * 33, 774, tier="Host", w=28, h=31)
    f.transfer_band(250, 790, 303, h=20)
    f.operation_label(276, 771, "D2H", size=19)
    f.centered(188, 838, "HBM", size=22)
    f.centered(358, 838, "Host", size=22)
    f.centered(272, 882, f.tr("stable FULL", "稳定 FULL"), size=22)
    f.token(87, 924, "C1", TEAL_PALE, r=20, size=18)
    f.check(109, 894, color=INK, size=17)
    for index in range(3):
        f.storage_page(140 + index * 33, 908, tier="HBM",
                       owner=GRAY if index == 0 else TEAL_PALE, w=28, h=31)
    f.lock(146, 943, scale=.5)
    f.text(140, 991, f.tr("shared", "共享"), 21)
    f.arrow(((246, 924), (302, 924)), color=ACTION)
    f.strip(314, 908, [PAPER] * 2, w=28, h=31)
    f.text(64, 1042, f.tr("RETURN: retire private KV", "RETURN：释放私有 KV"), 22)

    f.panel(498, 670, 430, 434, f.tr("D3  Prefetch + admission", "D3  预取与准入"),
            color=COORDINATION_HEADER, fill=PAPER, border=CONTROL_EDGE, size=26)
    f.token(536, 762, "P", PINK, r=22, size=22)
    f.clock(605, 762, 20)
    f.interval(653, 855, 762, color=INK, center=782)
    f.strip(526, 825, [PAPER] * 3 + [TEAL_PALE] * 2 + [BACKPLATE] * 2,
            w=45, h=34, gap=7)
    f.text(526, 898, f.tr("free / cold / growth", "空闲 / 冷页 / 增长"), 23)
    f.math(523, 953, [("c", None, "*"), " / ΔFULL"], size=29)
    f.diamond(816, 943, 150, 57, PAPER, INK)
    f.math(816, 953, [("τ",), " ≤ ", ("H",)], size=29, anchor="middle")
    f.text(523, 1004, "PREPARE", 22, INK, 700)
    f.text(733, 1004, "PREFETCH", 22, INK, 700)
    f.storage_page(526, 1035, tier="HBM", w=39, h=35)
    f.transfer_band(578, 1052, 632, h=22)
    f.storage_page(645, 1035, tier="Host", w=39, h=35)
    f.storage_page(736, 1035, tier="Host", w=39, h=35)
    f.transfer_band(788, 1052, 842, h=22)
    f.storage_page(855, 1035, tier="HBM", w=39, h=35)
    f.arrow(((247, 644), (247, 667)), color=BLUE, dashed=True)
    f.arrow(((706, 644), (706, 667)), color=BLUE, dashed=True)
    f.arrow(((481, 938), (495, 938)), color=ACTION)

    pool_width, pool_height = 764, 200
    for tier, yy, fills in (
        ("HBM", 350, [BACKPLATE] + [PINK] * 3 + [AMBER_PALE] * 3 + [PAPER] * 3),
        ("Host", 730, [BACKPLATE] + [PINK] * 5 + [AMBER_PALE] * 2 + [PAPER] * 2),
    ):
        f.memory_pool(986, yy, pool_width, pool_height, f"{tier} pool", tier)
        f.text(1005, yy + 92, "FULL", 24, INK, 700)
        for index, owner in enumerate(fills):
            if owner == PAPER:
                f.page(1125 + index * 59, yy + 65, fill=PAPER,
                       stroke=STORAGE_STYLES[tier]["edge"], w=52, h=39)
            else:
                f.storage_page(1125 + index * 59, yy + 65, tier=tier,
                               owner=GRAY if owner == BACKPLATE else owner,
                               w=52, h=39)
        f.text(1005, yy + 164, "Mamba", 23, INK, 700)
        for index, (fill, label) in enumerate((
            (PINK, "P"), (PAPER, "C1"), (AMBER_PALE, "C2"), (PAPER, ""),
        )):
            xx = 1128 + index * 154
            f.storage_state(xx, yy + 130, tier=tier, owner=fill,
                            scale=.6, exists=fill != PAPER)
            if label:
                f.text(xx + 61, yy + 157, label, 23)

    f.panel(986, 564, 330, 146, f.tr("Prefix index", "前缀索引"),
            color=BACKPLATE, fill=GRAY_PALE, border=GRAY, size=24, band=32)
    for xx, yy, label, fill in (
        (1006, 648, "p0", BACKPLATE), (1114, 613, "P", PINK),
        (1114, 674, "C2", AMBER_PALE), (1227, 613, "c*", PINK),
    ):
        f.page(xx, yy, label, fill=fill, w=54, h=32, size=22)
    f.line(1063, 661, 1110, 629, INK, 1.7)
    f.line(1063, 671, 1110, 690, INK, 1.7)
    f.arrow(((1173, 629), (1223, 629)), color=ACTION)

    f.centered(1544, 592, "HiCache", size=27, bold=True)
    f.transfer_engine(1474, 626, 140, 66)
    f.transfer_curve((1395, 552), (1344, 610), (1344, 670), (1395, 728))
    f.transfer_curve((1694, 728), (1741, 670), (1741, 610), (1694, 552))
    f.operation_label(1420, 603, "D2H", size=25)
    f.operation_label(1656, 603, "H2D", size=25)
    f.arrow(((1285, 629), (1326, 629), (1326, 657), (1461, 657)), color=ACTION)
    f.arrow(((930, 1025), (954, 1025), (954, 722), (1544, 722), (1544, 698)),
            color=ACTION)

    f.panel(986, 946, 764, 148, f.tr("Native admission", "原生准入"),
            color=DATA_PALE, fill=PAPER, border=DATA_EDGE, size=26, band=37)
    for yy, label, fill in ((999, "P", PINK), (1031, "hit", TEAL_PALE), (1063, "old", BACKPLATE)):
        f.ticket(1003, yy, label, fill, w=208, h=25, resident=label != "old")
    f.arrow(((1217, 1012), (1250, 1012), (1250, 1043), (1294, 1043)), color=ACTION)
    f.arrow(((1217, 1043), (1294, 1043)))
    f.arrow(((1217, 1076), (1250, 1076), (1250, 1055), (1294, 1055)),
            color=BLUE, dashed=True)
    f.text(1305, 1053, "KV ∧ slot", 29, INK, 700)
    f.arrow(((1466, 1043), (1562, 1043)))
    f.gpu(1580, 1000, 146, 87, running=[PINK, TEAL_PALE, AMBER_PALE, PINK, PAPER, PAPER])
    f.arrow(((930, 1076), (982, 1076)), color=ACTION)
    f.design_footer()
    return f


def prediction_and_overlap(language: str) -> PaperFigure:
    f = PaperFigure("02_prediction_and_overlap", "Execution-demand Inference", 660, language, width=1600)
    f.mechanism_title("D1", "Infer the next agent from completion signals",
                      "从完成信号推断近期执行需求")
    f.rect(657, 176, 230, 295, CONTROL_PALE, CONTROL_EDGE, 1.6, 8)

    f.token(64, 218, "P", PINK, r=25)
    f.token(177, 139, "C1", TEAL_PALE, r=25)
    f.token(177, 299, "C2", AMBER_PALE, r=25)
    f.check(207, 110, color=INK, size=19)
    f.diamond(299, 218, 87, 69, PAPER, INK)
    f.centered(299, 227, "ALL", size=25, bold=True)
    f.arrow(((91, 200), (149, 157)))
    f.arrow(((91, 237), (149, 282)))
    f.curve_arrow((205, 139), (244, 139), (241, 181), (257, 198), color=GRAY)
    f.curve_arrow((205, 299), (244, 299), (241, 254), (257, 238), color=BLUE, dashed=True)
    f.text(43, 373, "pending = {C2}", 23, INK, 700)
    f.arrow(((345, 218), (357, 218), (357, 102), (1395, 102), (1395, 164)),
            color=BLUE, dashed=True)
    f.step(710, 102, 1, "Bind the single pending child to its waiting parent")
    f.centered(920, 91, f.tr("bind the waiter", "确定等待对象"), size=23)
    f.token(1395, 196, "P", PINK, r=28)

    f.document(385, 246, 133, 119, color=AMBER_PALE,
               labels=(f.tr("… evidence", "… 证据"), f.tr("… checks", "… 验证"),
                       f.tr("… summary", "… 总结")))
    f.arrow(((205, 299), (379, 299)))
    f.polygon(((386, 397), (516, 397), (516, 434), (453, 434),
               (437, 447), (437, 434), (386, 434)), AMBER_PALE, INK)
    f.centered(451, 424, "NOTICE", size=21, bold=True)
    f.arrow(((521, 303), (594, 303), (594, 251), (661, 251)))
    f.arrow(((594, 303), (594, 365), (661, 365)))
    f.arrow(((520, 414), (554, 414), (554, 303), (575, 303)))
    f.step(574, 303, 2, "Infer final-phase confidence and conditional work remaining")
    f.centered(572, 210, f.tr("infer", "推断"), size=24, bold=True)
    f.math(766, 218, [("p", "final")], size=31, anchor="middle")
    f.rect(680, 237, 175, 26, PAPER, RULE, 1.2, 0)
    f.rect(680, 237, 125, 26, BLUE_PALE, INK, 1, 0)
    f.interval(680, 855, 365, color=INK, center=776)
    f.math(768, 409, [("W", "lo"), " … ", ("W", "hi")], size=29, anchor="middle")
    f.centered(768, 450, f.tr("tokens left if final", "终态剩余 token"), size=23)
    f.arrow(((863, 250), (930, 250), (930, 321)), color=ACTION)
    f.arrow(((864, 365), (930, 365), (930, 303), (1034, 303)), color=ACTION)
    f.step(982, 303, 3, "Advance remaining work using observed token progress and rate")
    f.centered(1015, 226, f.tr("roll forward", "滚动更新"), size=24, bold=True)
    f.clock(1065, 303, 24)
    f.arrow(((1095, 303), (1296, 303)), color=BLUE, dashed=True)
    f.math(1395, 400, [("τ",), " ≈ ", ("W",), " / ", ("v", "eff")],
           size=32, anchor="middle")
    for index, width in enumerate((25, 39, 58)):
        f.rect(1012 + index * 78, 420, width, 15, BLUE_PALE, INK, .9, 0)
    f.centered(1138, 481, f.tr("observed service", "已观测服务进度"), size=22)
    f.arrow(((1065, 413), (1065, 335)), color=BLUE, dashed=True)
    f.interval(1310, 1480, 303, color=INK, center=1405)
    f.line(1395, 227, 1395, 293, GRAY, 1.5, "5 4")
    f.text(1310, 347, "RETURN", 22)

    f.terminal(52, 520, 88, 62)
    f.clock(190, 550, 25)
    f.centered(134, 624, f.tr("tool + elapsed", "工具 + 已耗时"), size=24)
    f.arrow(((221, 550), (283, 550)))
    f.rect(400, 504, 122, 83, TEAL_PALE, "none", 0, 0)
    f.path(((293, 586), (342, 586), (342, 569), (400, 569),
            (400, 539), (457, 539), (457, 511), (522, 511)), INK, 2.5)
    f.line(400, 502, 400, 590, ACTION, 1.5, "5 4")
    f.centered(410, 624, f.tr("history | still running", "历史 | 仍未结束"), size=23)
    f.arrow(((535, 550), (735, 550)), color=ACTION)
    f.clock(770, 550, 24)
    f.math(770, 605, [("τ", "tool")], size=31, anchor="middle")
    f.arrow(((800, 550), (1296, 550)), color=BLUE, dashed=True)
    f.token(1395, 481, "A", TEAL_PALE, r=27)
    f.line(1395, 511, 1395, 540, GRAY, 1.5, "5 4")
    f.interval(1310, 1480, 550, color=INK, center=1405)
    f.centered(1395, 615, "TOOL_END", size=22)
    f.arrow(((1494, 303), (1520, 303), (1520, 428), (1542, 428)),
            color=BLUE, dashed=True)
    f.arrow(((1494, 550), (1520, 550), (1520, 428), (1542, 428)),
            color=BLUE, dashed=True)
    f.badge(1570, 428, "D3", fill=BACKPLATE, r=26)
    f.step(1520, 428, 4, "Pass candidate identity and timing to prefetch/admission")
    return f


def pages_and_handoff(language: str) -> PaperFigure:
    f = PaperFigure("03_pages_and_capacity_handoff", "Workflow-aware KV Lifecycle", 640, language, width=1600)
    f.mechanism_title("D2", "Back up stable KV; reclaim by workflow state",
                      "按工作流状态备份、保留与释放 KV")
    f.storage_band(210, 104, 880, 98, "HBM")
    f.storage_band(210, 298, 880, 98, "Host")
    f.storage_band(320, 450, 770, 99, "HBM")
    f.rect(1190, 104, 218, 477, GRAY_PALE, GRAY, 1.7, 8)
    f.token(85, 220, "P", PINK, r=28)
    f.centered(85, 279, f.tr("waiting", "等待中"), size=25)
    f.text(218, 160, "HBM", 27, INK, 700)
    f.text(218, 345, "Host", 27, INK, 700)
    for i in range(6):
        xx = 355 + i * 88
        f.storage_page(xx, 128, f"p{i}", tier="HBM", owner=GRAY if i == 0 else PINK,
                       w=64, h=46, size=24)
        f.storage_page(xx, 313, f"p{i}", tier="Host", owner=GRAY if i == 0 else PINK,
                       missing=i >= 4, w=64, h=46, size=24)
        if i >= 4:
            f.arrow(((xx + 32, 180), (xx + 32, 306)), color=AMBER)
    f.lock(367, 183, scale=.6)
    f.text(348, 236, f.tr("shared: keep", "共享：保留"), 22)
    f.bracket(531, 683, 112, color=ACTION)
    f.curve_arrow((607, 111), (850, 111), (1090, 150), (1215, 150), color=ACTION, width=6)
    f.step(965, 133, 2, "Reclaim cold FULL pages only after a stable backup exists")
    f.centered(1071, 89, f.tr("cold + backed", "冷页且已备份"), size=24)
    for i in range(2):
        f.page(1231 + i * 62, 128, fill=PAPER, w=45, h=46)
    f.centered(1293, 215, f.tr("+ capacity", "+ 可用容量"), size=24)
    f.centered(582, 260, f.tr("increment only", "仅传增量"), size=24)
    f.operation_label(583, 295, "D2H", size=25)
    f.step(827, 250, 1, "PREPARE_HOST copies missing stable FULL, not mutable Mamba")
    f.bracket(355, 859, 383, color=ACTION)
    f.math(632, 434, [("c", None, "*"), " : stable FULL"], size=28, anchor="middle")
    f.storage_state(942, 312, tier="Host", scale=.75)
    f.math(980, 434, [("s",), "(", ("c", None, "*"), ")"], size=29, anchor="middle")
    f.text(934, 275, f.tr("match if needed", "需要时匹配"), 23)

    f.token(86, 497, "C1", TEAL_PALE, r=27)
    f.check(119, 464, color=INK, size=19)
    f.text(135, 506, "RETURN", 25, INK, 700)
    f.arrow(((260, 497), (340, 497)))
    for i in range(4):
        f.storage_page(355 + i * 88, 474, f"p{i}", tier="HBM",
                       owner=GRAY if i == 0 else TEAL_PALE, w=64, h=46, size=24)
    f.lock(373, 530, scale=.6)
    f.bracket(443, 683, 537, color=ACTION)
    f.centered(566, 574, f.tr("unused private suffix", "无用私有后缀"), size=24)
    f.arrow(((700, 497), (1208, 497)), color=ACTION, width=6)
    f.step(965, 497, 3, "Retire unreferenced terminal-child private KV without D2H")
    f.centered(948, 530, f.tr("retire, no D2H", "直接释放，无须 D2H"), size=24)
    for i in range(3):
        f.page(1231 + i * 55, 474, fill=PAPER, w=41, h=46)
    f.centered(1304, 559, f.tr("+ capacity", "+ 可用容量"), size=24)
    f.token(86, 597, "C2", AMBER_PALE, r=25)
    f.lock(133, 579, scale=.7)
    f.text(180, 606, f.tr("running: keep live KV", "执行中：保留活跃 KV"), 25)

    f.path(((859, 383), (1035, 383), (1035, 337)), ACTION, 4.8)
    f.path(((1007, 337), (1446, 337)), ACTION, 4.8)
    f.path(((1359, 150), (1433, 150), (1433, 337)), ACTION, 4.8)
    f.path(((1393, 497), (1433, 497), (1433, 337)), ACTION, 4.8)
    f.arrow(((1446, 337), (1510, 337)), color=ACTION)
    f.step(1446, 337, 4, "Expose checkpoint, missing pages, matching state and released capacity")
    f.centered(1221, 306, f.tr("restore view", "恢复视图"), size=25, bold=True)
    f.badge(1547, 337, "D3", fill=BACKPLATE, r=28)
    return f


def prefetch_and_admission(language: str) -> PaperFigure:
    f = PaperFigure("04_prefetch_and_admission", "Prefetch and Admission Coordination", 620, language, width=1600)
    f.mechanism_title("D3", "Restore early, then admit for first use",
                      "提前恢复，协同准入与首次消费")
    f.storage_band(132, 222, 442, 135, "Host")
    f.storage_band(1097, 222, 478, 135, "HBM")
    f.rect(712, 180, 280, 65, CONTROL_PALE, CONTROL_EDGE, 1.6, 8)
    f.rect(132, 423, 442, 145, GRAY_PALE, GRAY, 1.6, 8)
    f.rect(1097, 384, 478, 202, DATA_PALE, DATA_EDGE, 1.8, 8)

    f.token(72, 133, "C2", AMBER_PALE, r=28)
    f.line(117, 133, 1000, 133, RULE, 1.3)
    f.rect(492, 113, 508, 40, AMBER_PALE, INK, 1.5, 0)
    f.centered(746, 141, f.tr("final report", "最终报告生成"), size=26)
    f.circle(420, 133, 7, BLUE, BLUE, 1.4)
    f.text(378, 98, "NOTICE", 24, INK, 700)
    f.line(1000, 92, 1000, 568, GRAY, 1.7, "7 6")
    f.text(1021, 139, "RETURN / JOIN", 25, INK, 700)
    f.badge(364, 190, "D1", fill=BLUE_PALE, r=23)
    f.arrow(((420, 157), (420, 166), (733, 166), (733, 276)),
            color=BLUE, dashed=True)
    f.step(576, 166, 1, "Use the rolling completion hint to choose the transfer start")
    f.text(779, 221, f.tr("time restore", "决定启动时机"), 25)
    f.clock(963, 213, 18)

    f.token(72, 300, "P", PINK, r=28)
    f.text(148, 253, "Host", 27, INK, 700)
    f.storage_page(259, 277, "p2", tier="Host", w=65, h=46, size=24)
    f.storage_page(345, 277, "p3", tier="Host", w=65, h=46, size=24)
    f.bracket(259, 410, 344, color=ACTION)
    f.math(334, 388, ["ΔFULL"], size=29, anchor="middle")
    f.storage_state(463, 278, tier="Host", scale=.75)
    f.math(493, 386, [("s",), "(", ("c", None, "*"), ")"], size=28, anchor="middle")
    f.centered(489, 260, f.tr("if needed", "必要时"), size=22)
    f.path(((410, 344), (550, 344), (550, 300)), AMBER, 5)
    f.arrow(((530, 300), (671, 300)), color=AMBER)
    f.activity_bar(679, 280, 392, 40)
    f.operation_label(924, 309, "H2D", size=28)
    f.step(825, 300, 3, "Restore only the missing FULL pages and required matching state")
    f.text(1111, 253, "HBM", 27, INK, 700)
    for i in range(6):
        f.storage_page(1112 + i * 72, 278, f"p{i}", tier="HBM",
                       owner=GRAY if i == 0 else PINK, w=56, h=45, size=23)
    f.bracket(1256, 1384, 267, color=ACTION)
    f.storage_state(1444, 231, tier="HBM", scale=.6)
    f.math(1540, 258, [("c", None, "*")], size=29, anchor="middle")
    f.lock(1371, 328, scale=.65)
    f.text(1410, 350, f.tr("bounded", "有界驻留"), 22)

    f.badge(73, 462, "D2", fill=TEAL_PALE, r=23)
    f.strip(142, 441, [PAPER, PAPER, TEAL_PALE, BACKPLATE], w=49, h=40, gap=12)
    f.lock(329, 413, scale=.6)
    f.text(148, 529, f.tr("free / reclaimable / growth", "空闲 / 可回收 / 增长预留"), 24)
    f.arrow(((392, 460), (600, 460), (600, 340), (692, 340), (692, 324)), color=ACTION)
    f.step(600, 405, 2, "Select missing pages within per-pool capacity and growth reserve")
    f.text(681, 449, f.tr("select within budget", "按容量预算选择"), 25)

    f.text(1111, 414, f.tr("native queue", "原生队列"), 25, INK, 700)
    for yy, label, fill in ((427, "P", PINK), (471, "hit", TEAL_PALE), (515, "old", BACKPLATE)):
        f.ticket(1112, yy, label, fill, w=217, h=32,
                 resident=label != "old", protected=False)
    f.arrow(((1030, 531), (1070, 531), (1070, 443), (1108, 443)), color=BLUE, dashed=True)
    f.centered(983, 600, "JOIN + request", size=23)
    f.gpu(1443, 420, 116, 99, running=[PINK, TEAL_PALE, PINK, PAPER, PAPER, PAPER])
    f.arrow(((1336, 443), (1430, 443)), color=FLOW, width=6)
    f.step(1383, 443, 4, "Admit restored runnable work under native data and slot dependencies")
    f.text(1378, 570, f.tr("data + slot", "数据 + slot"), 23)
    f.text(1111, 570, f.tr("bounded + aging", "有界优先 + 老化"), 23)
    return f


def export_collection(figures, directory: Path):
    destination = directory / "beliefkv_figures_bilingual.pdf"
    subprocess.run(
        ["pdfunite", *(str(directory / figure.filename.replace(".svg", ".pdf"))
                       for figure in figures), str(destination)],
        check=True,
    )


def slide_notes(figure: PaperFigure):
    if figure.filename.startswith("01_"):
        return figure.tr(
            "Architecture baseline: 2026-10-15 development implementation, not the frozen "
            "code of every running experiment. RCCG receives runtime events; GraphDelta "
            "updates affected wait groups. Design IDs are consistent across all four figures: "
            "D1 execution-demand inference, D2 workflow-aware KV lifecycle, D3 prefetch/"
            "admission coordination. The shared footer shows D1 hints to D3, D2 views to D3, "
            "D3 physical actions to D2, and actual service/progress feedback to D1. "
            "The running example is parent P and children C1/C2. Text/phase/work inference runs in an independent "
            "CPU process. Tool history supplies separate conditional timing. "
            "The wait-completion panel has two distinct paths: delivered child text, "
            "notice and generated progress feed final-stage confidence and conditional "
            "remaining tokens in the async CPU worker; runtime converts work to a "
            "RETURN timing hint using observed effective rate. The single pending child "
            "is associated with its waiting parent by RCCG. Tool command history and "
            "elapsed age instead produce a survival-conditioned TOOL_END hint for the "
            "tool-owning agent A. Dashed candidate arrows do not mark either agent ready. "
            "The panel sends identity/timing hints to cache and admission planning; "
            "neither path independently authorizes transfers or execution. "
            "NativeAdmissionRuntime combines causal identity, checkpoint, timing, capacity "
            "and next-batch observations. Waiting migrations can run after forward dispatch "
            "under the supported overlap configuration. The native waiting queue, "
            "PrefillAdder, allocation and data dependencies remain authoritative. "
            "Dashed arrows show feedback, not a claim that all control work is asynchronous. "
            "Explicit completion notices are harness protocol and must be identical "
            "across comparison arms. Root workflows can spawn again after JOIN; the current "
            "child harness does not expose nested task delegation. A tool completion "
            "resumes its child; only a child RETURN contributes to JOIN completion. "
            "Rounded groups encode hierarchy on an aligned grid. The control plane uses "
            "blue borders/background and the data plane green. HBM and Host share tier "
            "geometry and black outlines but use distinct purple/gold headers and page "
            "bodies with neutral/peach backgrounds. All explicit colors are sampled from "
            "the user reference, with provenance in the manifest. Thin owner stripes "
            "and page labels remain unchanged across copies, "
            "separating residency from agent identity. The prefix index is a compact "
            "neutral group, not a third storage tier. Native admission has a separate "
            "service group. "
            "Location and solid/dashed page borders indicate residency. Triangular arrows "
            "show execution, notched arrows show runtime control, broad curved ribbons "
            "show KV movement, and dashed open curves show feedback. HiCache/DMA denotes "
            "the transfer subsystem; D2H and H2D label actions, not separate components. Hexagons "
            "marked S identify Mamba checkpoint states, separate from FULL page rectangles. "
            "NOTICE is an observed completion announcement, not ground-truth terminal phase. "
            "The control plane combines checkpoints, measured timing, native capacity and "
            "candidate order; physical FULL pages and Mamba slots remain separate. "
            "Sources: causal_graph.py; semantic_report_worker.py; sglang_v0520_runtime.py "
            "(scheduler_step, dispatch_after_forward, plan_native_prefill); staging patch.",
            "架构基线：2026-10-15 当前开发实现，不代表所有运行实验的冻结代码。"
            "RCCG 接收运行时事件，GraphDelta 增量更新受影响的等待成员组。"
            "四图设计编号一致：D1 执行需求推断，D2 工作流感知的 KV 生命周期，"
            "D3 准入与预取协同调度。页脚共同表达 D1 向 D3 提供候选与时机、"
            "D2 向 D3 提供物理状态、D3 更新 D2 的缓存生命周期、实际服务进度反馈到 D1。"
            "同一示例中的 P 为 parent，C1/C2 为其 children。"
            "文本阶段/工作量推理在独立 CPU 进程执行，工具历史提供条件时机提示。"
            "等待结束预测框分成两路：child 已送达正文、通知与生成进度进入异步 CPU "
            "模型，输出最终报告阶段置信度及条件剩余 token；runtime 再结合观测有效速率"
            "给出 RETURN 时机提示。RCCG 将唯一未返回 child 对应到等待中的 parent。"
            "工具一路则用命令历史与已耗时，在仍未完成的条件下估计 TOOL_END，"
            "对应到工具所属 agent A。虚线只表示候选恢复需求，不把 P/A 标为 ready。"
            "该框向缓存与准入规划传递身份和时机提示，两路都不独立授权迁移或执行。"
            "NativeAdmissionRuntime 联合因果身份、checkpoint、时机、容量和下一批观测。"
            "受支持的 overlap 配置下，等待迁移规划安排在 forward 发射后。"
            "原生可见等待队列、PrefillAdder、分配器和数据依赖保持权威。"
            "虚线表达观测反馈，不表示所有控制工作均异步。"
            "显式完成通知属于 harness 协议，应在所有对照组保持一致。"
            "root 可以在 JOIN 后再次 spawn；当前 child harness 不提供嵌套 task 委派。"
            "工具完成只让 child 恢复执行，只有实际 child RETURN 才计入 JOIN 完成。"
            "圆角分组与对齐网格表达层级。控制面为蓝边浅蓝底，数据面为绿边浅绿底。"
            "HBM/Host 使用对等层级的几何形状和黑色边框，标题与页体分别采用紫色和黄色，"
            "背景分别为浅灰和浅桃色。全部显式颜色取自用户原图，取色记录保存在 manifest。"
            "相同页号和细 owner 色条跨副本保持一致，将存储位置与 agent 归属分开。"
            "前缀索引为紧凑中性分组，不是第三个存储池；原生准入有独立服务分组。"
            "驻留由位置及实虚边框表达。三角箭头为执行、凹尾箭头为控制、"
            "宽曲线带为 KV 迁移、开口虚线曲线为反馈。HiCache/DMA 表示传输子系统；"
            "D2H/H2D 标注传输边，不是独立设备。六边形 S 表示 Mamba checkpoint，"
            "与矩形 FULL 页分开。"
            "NOTICE 是已观测到的收尾通知，不是真实终态标签。控制面结合 checkpoint、"
            "实测时机、原生容量和候选顺序；数据面的 FULL 页与 Mamba 槽分别核算。"
            "代码：causal_graph.py、semantic_report_worker.py、sglang_v0520_runtime.py "
            "的 scheduler_step / dispatch_after_forward / plan_native_prefill，以及 staging patch。",
        )
    if figure.filename.startswith("02_"):
        return figure.tr(
            "The numbered circles annotate execution arrows. 1: bind the single pending "
            "child C2 to waiting parent P; C1 has already returned. 2: infer final-phase "
            "confidence and conditional remaining-token bounds from delivered text, notice "
            "and generated progress in an independent CPU worker. 3: roll the current work "
            "estimate forward and convert it with observed effective service. W in the "
            "compact formula is current remaining work after subtracting only progress "
            "since its forecast snapshot. The 0.5/2/5-second effective-rate windows include "
            "past service gaps, never future queueing. The progress arrow supplies rolling updates. "
            "4: pass candidate identity and timing to D3. The separate tool path conditions "
            "history on the tool still running and identifies owning agent A; it does not "
            "share the content model. Current live tool transfer handles at most one active "
            "call. The drawing abstracts encoder internals, and all confidence/work/time "
            "marks are schematic. Forecasts, NOTICE and EOS do not satisfy JOIN or authorize "
            "H2D. The current JOIN content path requires one pending child; work bounds are "
            "not a calibrated multi-child JOIN wall-clock distribution. Sources: "
            "causal_graph.py; child_semantic_work.py; conditional_work.py; "
            "project_tool_history.py; semantic_report_worker.py; runtime semantic/tool paths.",
            "黑色圆圈标在执行箭头上。①将唯一未返回的 C2 绑定到等待中的 P，C1 已返回。"
            "②独立 CPU worker 根据已送达正文、通知与生成进度推断最终阶段置信度和"
            "条件剩余 token。③扣除预测快照之后新增的进度，滚动更新当前剩余工作，再"
            "结合观测有效速率换算时机。简式中的 W 已是更新后的剩余工作，不能重复扣除。"
            "0.5/2/5 秒速率包含已发生的服务间隙，未来排队不是在线输入；进度箭头表示滚动更新。"
            "④将候选身份与时机交给 D3。工具路径独立根据历史、已耗时及仍未结束这一条件"
            "估计 TOOL_END，并绑定工具所属 agent A；不经过内容模型。在线工具迁移路径"
            "至多处理一个活跃调用。主图省略编码器内部结构，分数和区间均为机制示意。"
            "NOTICE、EOS 与预测不直接满足 JOIN，也不授权 H2D；当前内容 JOIN 路径要求"
            "唯一未返回 child，工作区间不代表已校准的多 child JOIN 墙钟分布。"
            "代码：causal_graph.py、"
            "child_semantic_work.py、conditional_work.py、project_tool_history.py、"
            "semantic_report_worker.py 与 runtime 的语义/工具路径。",
        )
    if figure.filename.startswith("03_"):
        return figure.tr(
            "Execution-arrow numbers identify typical operations, not mandatory global "
            "ordering. 1: copy only missing stable FULL pages p4/p5 to Host; p0..p3 already "
            "have backups. PREPARE_HOST itself neither frees HBM nor speculatively backs "
            "mutable Mamba. 2: under pressure, reclaim already-backed cold p2/p3 in HBM. "
            "3: C1 RETURN retires its unused private suffix p1..p3 without D2H; shared p0 "
            "and C2 live KV remain. Reclamation requires no active references, shared "
            "dependencies or in-flight transfer hazards. Empty page symbols show released "
            "capacity, not copied child data or parent-exclusive reservations. 4: expose "
            "the safe-input checkpoint c*, missing FULL, matching necessary state and "
            "actual free capacity to D3. c* ends at p5; generated suffixes are excluded. "
            "The state hexagon is a separate matching-state lookup, not another PREPARE "
            "destination; newer decoded Mamba need not match c*. Short waits can simply "
            "retain HBM. Physical pages and FULL/Mamba pools remain native-managed. Sources: "
            "runtime waiting prepare/terminal reclaim; sglang_v0520_physical.py; staging patch.",
            "箭头序号对应典型操作，不表示全局强制执行顺序。①只补传 Host 中缺失的稳定"
            "FULL p4/p5，p0..p3 已有副本；PREPARE_HOST 本身不释放 HBM，也不投机备份"
            "持续更新的 Mamba。②有容量压力时回收 HBM 已备份的冷页 p2/p3。③C1 RETURN"
            "后直接释放已无用途的私有 p1..p3，无须 D2H；共享 p0 与 C2 活跃 KV 保留。"
            "释放必须满足无活跃引用、共享依赖及在途传输风险等原生条件。空页表示释放的"
            "容量，不表示 child 数据复制给 parent 或为其独占预留。④将安全输入 checkpoint "
            "c*、缺失 FULL、必要匹配状态及实际容量提供给 D3。c* 结束于 p5，生成后缀不"
            "作为恢复目标。状态六边形表示独立的匹配状态查找，不是 PREPARE 的另一个"
            "传输目标；最新 decode 状态未必匹配 c*。短等待可以直接保留 HBM。"
            "物理页及 FULL/Mamba 分池继续由原生系统管理。代码：runtime 等待备份/终态回收、"
            "sglang_v0520_physical.py 与原生缓存 staging patch。",
        )
    return figure.tr(
        "Execution-arrow steps coordinate inference and native service. 1: use D1's "
        "rolling hint and measured direction/shape/size-matched transfer cost to time "
        "restoration. 2: use D2's free/reclaimable capacity with next-input growth reserve, "
        "accounted separately for FULL and Mamba. 3: H2D copies only missing p2/p3 and "
        "necessary state matching c*; other pages in the right-hand prefix were resident. "
        "Two pages are an example, not a fixed cap. Its start before and completion after "
        "the actual C2 RETURN line illustrate useful partial overlap; no empirical "
        "timing scale is claimed. C1 has already returned, as in D1/D2. 4: after actual "
        "JOIN and request arrival, admit restored runnable P with a bounded boost, "
        "while retaining HBM-hit preference, aging and batch fill. Native PrefillAdder, "
        "allocator, layer dependencies, required Mamba and slots remain authoritative. "
        "The lock denotes bounded protection pending first use, not an indefinite pin. "
        "Ledger ACK manages lifecycle but is not a universal execution barrier. Other "
        "tails/dependencies use native recovery; failures or expiry fall back natively. "
        "Control paths reuse current-cycle observations and remain lightweight without "
        "opportunities. No throughput or JCT gain is implied by the schematic. Sources: "
        "native_transfer_policy.py; native_transfer_service.py; runtime "
        "_wait_prefetch_opportunity, defer_prefill_for_prefetch, plan_native_prefill.",
        "箭头序号表达推断与原生服务的协同。①根据 D1 的滚动提示，以及同方向/形态/大小"
        "的实测传输耗时选择启动时机。②根据 D2 的空闲、可回收容量及下一批输入增长"
        "预留选择缺页，FULL/Mamba 分池核算。③只 H2D 缺失的 p2/p3 及匹配 c* 的必要"
        "状态，右侧其余页已驻留。两页是示例，不是固定上限。传输在真实 C2 RETURN 前"
        "开始、之后完成，表示有用的部分重叠，不代表实测时间。C1 此前已返回，与 D1/D2"
        "一致。④真实 JOIN 完成且请求提交后，对已恢复且可执行的 P 有界提升，同时保留"
        "HBM 命中倾向、老化和 batch 填充。原生 PrefillAdder、分配器、逐层数据依赖、"
        "必要 Mamba 与 slot 保持决定权。锁表示待首次消费期间的有界保护，不是永久固定。"
        "账本 ACK 管理生命周期，不是所有执行都要等待的统一屏障；其余尾部或依赖仍用"
        "原生恢复，失败、过期时原生回退。同轮观测复用和无机会轻量路径降低控制成本。"
        "示意图不声称已获得吞吐或 JCT 收益。代码：native_transfer_policy.py、native_transfer_service.py "
        "及 runtime 的 _wait_prefetch_opportunity、defer_prefill_for_prefetch、"
        "plan_native_prefill。",
    )


class PowerPointExporter:
    """Map SVG drawing primitives to editable DrawingML objects."""

    def __init__(self, language: str, *, width=WIDTH, height=HEIGHT):
        self.language = language
        self.presentation = Presentation()
        self.presentation.slide_width = Inches(14.4)
        self.presentation.slide_height = Inches(14.4 * height / width)
        self.unit = Inches(14.4 / width)
        self.point_scale = 14.4 * 72 / width
        self.presentation.core_properties.title = (
            "BeliefKV 系统架构图" if language == "zh" else "BeliefKV Architecture Figures"
        )
        self.presentation.core_properties.author = "BeliefKV"
        self.presentation.core_properties.subject = (
            f"Current development architecture as of {ARCHITECTURE_DATE}"
        )
        self.presentation.core_properties.keywords = (
            "BeliefKV, dynamic agent workflows, KV cache, native editable diagrams"
        )

    def length(self, pixels: float):
        return round(pixels * self.unit)

    @staticmethod
    def rgb(value: str):
        return RGBColor.from_string(value.lstrip("#"))

    @staticmethod
    def clear_shape_effects(shape):
        for effect in shape._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
        for effect in shape._element.xpath(".//a:effectLst|.//a:effectDag"):
            effect.getparent().remove(effect)
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        for properties in shape._element.xpath(".//a:rPr|.//a:defRPr|.//a:endParaRPr"):
            properties.append(OxmlElement("a:effectLst"))

    def clear_package_effects(self):
        namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
        removed = {
            "outerShdw", "innerShdw", "prstShdw", "glow", "reflection",
            "softEdge", "blur", "scene3d", "sp3d",
        }
        for part in self.presentation.part.package.iter_parts():
            if not part.content_type.endswith("+xml"):
                continue
            live_root = getattr(part, "_element", None)
            root = live_root if live_root is not None else etree.fromstring(part.blob)
            for element in list(root.iter()):
                name = etree.QName(element)
                if name.namespace != namespace:
                    continue
                if name.localname in ("effectLst", "effectDag"):
                    for child in list(element):
                        element.remove(child)
                elif name.localname in removed and element.getparent() is not None:
                    element.getparent().remove(element)
                elif name.localname == "effectRef":
                    element.set("idx", "0")
            if live_root is None:
                part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    def style(self, shape, attributes):
        fill = attributes.get("fill", "none")
        if shape.shape_type != MSO_SHAPE_TYPE.LINE:
            if fill == "none":
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = self.rgb(fill)
        stroke = attributes.get("stroke", "none")
        if stroke == "none":
            shape.line.fill.background()
            return
        shape.line.color.rgb = self.rgb(stroke)
        shape.line.width = Pt(float(attributes.get("stroke-width", 1)) * self.point_scale)
        if "stroke-dasharray" in attributes:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def add_arrowhead(self, slide, points, attributes):
        end = points[-1]
        previous = next((point for point in reversed(points[:-1]) if point != end), None)
        if previous is None:
            return
        length = math.hypot(end[0] - previous[0], end[1] - previous[1])
        ux, uy = (end[0] - previous[0]) / length, (end[1] - previous[1]) / length
        size = float(attributes.get("data-arrow-size", 16))
        kind = attributes.get("data-arrow-kind", "filled")
        base = (end[0] - size * ux, end[1] - size * uy)
        left = (base[0] - .5 * size * uy, base[1] + .5 * size * ux)
        right = (base[0] + .5 * size * uy, base[1] - .5 * size * ux)
        vertices = (
            [left, end, right] if kind == "open" else
            [end, left, right] if kind == "triangle" else
            [end, left, (end[0] - .75 * size * ux, end[1] - .75 * size * uy), right]
        )
        builder = slide.shapes.build_freeform(*vertices[0], scale=self.unit)
        builder.add_line_segments(vertices[1:], close=kind != "open")
        head = builder.convert_to_shape()
        head.name = f"arrowhead {len(slide.shapes)}"
        color = attributes.get("stroke", FLOW)
        self.style(head, {
            "fill": "none" if kind == "open" else color,
            "stroke": color if kind == "open" else "none",
            "stroke-width": str(size * 2.5 / 16),
        })
        self.clear_shape_effects(head)

    def add_bezier_shape(self, slide, commands):
        points = [point for _, part in commands for point in part]
        builder = slide.shapes.build_freeform(*points[0], scale=self.unit)
        builder.add_line_segments(points[1:], close=False)
        shape = builder.convert_to_shape()
        path = shape._element.xpath("p:spPr/a:custGeom/a:pathLst/a:path")[0]
        for child in list(path):
            path.remove(child)
        left = min(round(x) for x, _ in points)
        top = min(round(y) for _, y in points)
        for command, vertices in commands:
            operation = OxmlElement({
                "M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo", "Z": "a:close",
            }[command])
            for x, y in vertices:
                point = OxmlElement("a:pt")
                point.set("x", str(round(x) - left))
                point.set("y", str(round(y) - top))
                operation.append(point)
            path.append(operation)
        return shape

    def add_text(self, slide, element, figure, offset_y):
        value = "".join(element.itertext())
        size = float(element.get("font-size", 24))
        weight = int(element.get("font-weight", 400))
        family = element.get("font-family", figure.family)
        italic = element.get("font-style") == "italic"
        font = drawing_font(figure.language, size, weight, family, italic)
        width = figure.measure(value, size, weight, family, italic)[0]
        ascent, descent = (metric / 4 for metric in font.getmetrics())
        # Keep spare width for Office font substitution; alignment preserves SVG anchors.
        box_width = width + max(12, size * .3)
        x, y = float(element.get("x", 0)), float(element.get("y", 0)) + offset_y
        anchor = element.get("text-anchor", "start")
        left = x - (box_width / 2 if anchor == "middle" else box_width if anchor == "end" else 0)
        shape = slide.shapes.add_textbox(
            self.length(left), self.length(y - ascent),
            self.length(box_width), self.length(ascent + descent + 4),
        )
        shape.name = value[:100]
        frame = shape.text_frame
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.word_wrap = False
        paragraph = frame.paragraphs[0]
        paragraph.alignment = {
            "start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT,
        }[anchor]
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = value
        run.font.name = family
        run.font.size = Pt(size * self.point_scale)
        run.font.bold = weight >= 600
        run.font.italic = italic
        run.font.color.rgb = self.rgb(element.get("fill", INK))
        properties = run._r.get_or_add_rPr()
        properties.set("lang", "zh-CN" if figure.language == "zh" else "en-US")
        properties.set("spc", "0")
        east_asian = OxmlElement("a:ea")
        east_asian.set("typeface", family)
        properties.append(east_asian)
        self.clear_shape_effects(shape)

    def add_primitive(self, slide, element, figure, offset_y):
        tag = element.tag.rsplit("}", 1)[-1]
        attributes = dict(element.attrib)
        if tag == "text":
            self.add_text(slide, element, figure, offset_y)
            return
        if tag == "rect":
            if element.get("width") == "100%":
                return
            shape = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if float(element.get("rx", 0))
                else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                self.length(float(element.get("x", 0))),
                self.length(float(element.get("y", 0)) + offset_y),
                self.length(float(element.get("width", 0))),
                self.length(float(element.get("height", 0))),
            )
            if shape.adjustments:
                shape.adjustments[0] = (
                    float(element.get("rx", 0))
                    / min(float(element.get("width")), float(element.get("height")))
                )
        elif tag == "circle":
            radius = float(element.get("r", 0))
            shape = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL,
                self.length(float(element.get("cx", 0)) - radius),
                self.length(float(element.get("cy", 0)) - radius + offset_y),
                self.length(radius * 2), self.length(radius * 2),
            )
        elif tag == "line":
            points = [
                (float(element.get("x1", 0)), float(element.get("y1", 0)) + offset_y),
                (float(element.get("x2", 0)), float(element.get("y2", 0)) + offset_y),
            ]
            shape = slide.shapes.add_connector(
                MSO_CONNECTOR_TYPE.STRAIGHT,
                *(self.length(coordinate) for point in points for coordinate in point),
            )
        elif tag in ("path", "polygon"):
            if tag == "path":
                commands = vector_commands(element.get("d", ""), offset_y)
                points = [point for _, part in commands for point in part]
            else:
                numbers = re.findall(r"[-+]?(?:\d*\.\d+|\d+)", element.get("points", ""))
                points = [
                    (float(numbers[index]), float(numbers[index + 1]) + offset_y)
                    for index in range(0, len(numbers), 2)
                ]
                commands = []
            if any(command == "C" for command, _ in commands):
                shape = self.add_bezier_shape(slide, commands)
                tag = "curve"
            elif len(points) == 2:
                shape = slide.shapes.add_connector(
                    MSO_CONNECTOR_TYPE.STRAIGHT,
                    *(self.length(coordinate) for point in points for coordinate in point),
                )
            else:
                builder = slide.shapes.build_freeform(*points[0], scale=self.unit)
                builder.add_line_segments(points[1:], close=tag == "polygon")
                shape = builder.convert_to_shape()
        else:
            raise ValueError(f"Unsupported diagram element {tag}")
        shape.name = f"{tag} {len(slide.shapes)}"
        self.style(shape, attributes)
        self.clear_shape_effects(shape)
        if "marker-end" in attributes:
            self.add_arrowhead(slide, points, attributes)

    def add_figure(self, figure: PaperFigure, directory: Path):
        slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self.rgb(PAPER)
        self.unit = Inches(14.4 / figure.width)
        self.point_scale = 14.4 * 72 / figure.width
        offset_y = (self.presentation.slide_height / self.unit - figure.height) / 2
        root = ET.parse(directory / figure.filename).getroot()
        for element in root:
            if element.tag.rsplit("}", 1)[-1] in ("title", "defs", "style"):
                continue
            self.add_primitive(slide, element, figure, offset_y)
        slide.notes_slide.notes_text_frame.text = slide_notes(figure)

    def save(self, path: Path, *, expected_slides=4):
        self.clear_package_effects()
        self.presentation.save(path)
        reopened = Presentation(path)
        slides = []
        for slide in reopened.slides:
            pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
            if pictures:
                raise ValueError("Architecture slide contains a flattened picture")
            for shape in slide.shapes:
                if (
                    shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > reopened.slide_width
                    or shape.top + shape.height > reopened.slide_height
                ):
                    raise ValueError(f"PowerPoint object outside slide: {shape.name!r}")
            slides.append({
                "objects": len(slide.shapes),
                "text_objects": sum(shape.has_text_frame and bool(shape.text) for shape in slide.shapes),
                "picture_objects": len(pictures),
            })
        if len(slides) != expected_slides:
            raise ValueError(f"Expected {expected_slides} editable slides")
        return {"file": path.name, "language": self.language, "slides": slides}


def export_powerpoints(figures, directory: Path):
    outputs = []
    for language in ("zh", "en"):
        exporter = PowerPointExporter(language)
        for figure in figures:
            if figure.language == language:
                exporter.add_figure(figure, directory)
        outputs.append(exporter.save(directory / f"beliefkv_architecture_editable_{language}.pptx"))
    return outputs


def export_standalone_powerpoints(figures, directory: Path):
    outputs = []
    for figure in figures:
        if figure.filename.startswith("01_"):
            continue
        exporter = PowerPointExporter(
            figure.language, width=figure.width, height=figure.height,
        )
        exporter.add_figure(figure, directory)
        name = figure.filename.replace(".svg", "_editable.pptx")
        output = exporter.save(directory / name, expected_slides=1)
        output["figure"] = figure.filename.removesuffix(".svg")
        output["canvas"] = [figure.width, figure.height]
        outputs.append(output)
    return outputs


def preview(directory: Path, manifest: list[dict]):
    tile_width = 1020
    rows = []
    for row in manifest:
        if row["language"] != "en":
            continue
        with Image.open(directory / (row["stem"] + ".png")) as original:
            image = original.copy()
        image.thumbnail((tile_width - 48, 690), Image.Resampling.LANCZOS)
        tile = Image.new("RGB", (tile_width, image.height + 24), PAPER)
        tile.paste(image, (24, 12))
        rows.append(tile)
    sheet = Image.new("RGB", (tile_width, sum(row.height for row in rows) + 24), PAPER)
    position = 8
    for row in rows:
        sheet.paste(row, (0, position))
        position += row.height + 4
    sheet.save(directory / "preview.png")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    figures = [
        builder(language)
        for builder in (system_overview, prediction_and_overlap, pages_and_handoff,
                        prefetch_and_admission)
        for language in ("en", "zh")
    ]
    manifest = [figure.save_to(args.output) for figure in figures]
    export_collection(figures, args.output)
    powerpoints = export_powerpoints(figures, args.output)
    standalone_powerpoints = export_standalone_powerpoints(figures, args.output)
    preview(args.output, manifest)
    revision = subprocess.check_output(
        ["git", "-C", str(ROOT), "rev-parse", "HEAD"], text=True,
    ).strip()
    (args.output / "manifest.json").write_text(json.dumps({
        "source_revision": revision, "architecture_date": ARCHITECTURE_DATE,
        "source_revision_scope": (
            "latest development implementation, not every frozen experiment; "
            "source files and drawing generator are individually fingerprinted"
        ),
        "architecture_source_sha256": {
            name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest()
            for name in ARCHITECTURE_SOURCES
        },
        "generator_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "drawing_helpers_sha256": hashlib.sha256(
            (Path(__file__).parent / "generate_beliefkv_paper_figures.py").read_bytes(),
        ).hexdigest(),
        "schematic_not_empirical_performance": True,
        "figure_designs": [
            {"stem": "01_system_overview", "design": None, "role": "shared closed loop"},
            {"stem": "02_prediction_and_overlap", "design": "D1", "role": "execution demand"},
            {"stem": "03_pages_and_capacity_handoff", "design": "D2", "role": "KV lifecycle"},
            {"stem": "04_prefetch_and_admission", "design": "D3", "role": "prefetch/admission"},
        ],
        "shared_example": {
            "agents": ["P", "C1", "C2"],
            "checkpoint": "c* at p5; stable FULL p0..p5; D2 reclaims backed cold p2/p3",
            "sequence": [
                "P waits after SPAWN", "C1 RETURN frees unused private capacity",
                "C2 final report provides a forecast", "D3 restores P using D2 view",
                "C2 RETURN satisfies JOIN", "P request arrives and consumes restored KV",
            ],
            "feedback": "actual service, progress and completion update the loop",
        },
        "presentation_scope": (
            "one overview plus three compact algorithms with functional rounded groups; "
            "numbers annotate execution arrows; implementation details in slide notes"
        ),
        "visual_style": (
            "aligned rounded groups, distinct control/data planes and HBM/Host locations; "
            "independent owner stripes, numbered execution arrows and editable curves"
        ),
        "shadow_policy": (
            "no offset decorative layers; explicit shape/text resets; package-wide "
            "theme, master, layout, shadow, glow, reflection and 3D effects removed"
        ),
        "powerpoint_arrowheads": (
            "editable triangle, notched and open heads; native cubic Bezier paths, "
            "rounded control routes, closed curved transfer ribbons and wide block arrows"
        ),
        "color_policy": {
            "families": ["control blue", "data green", "HBM lavender", "Host gold", "neutral"],
            "mechanism_header_colors": [
                CONTROL_EDGE, DATA_EDGE, CONTROL_HEADER, LIFECYCLE_HEADER,
                COORDINATION_HEADER, BACKPLATE, DATA_PALE,
            ],
            "separate_hue_per_module": False,
            "accent_contrast": "distinct location bodies with stable thin owner stripes",
            "text_contrast_preserved": True,
            "overview_memory_tiers": (
                "HBM and Host share tier geometry and black outlines, with different background/page "
                "colors; blue control and green data planes; neutral prefix index"
            ),
            "explicit_colors_restricted_to_reference_samples": True,
        },
        "reference_palette": REFERENCE_PALETTE,
        "reference_palette_sha256": hashlib.sha256(REFERENCE_PALETTE_PATH.read_bytes()).hexdigest(),
        "typography": {
            "text_colors": [INK, PAPER],
            "page_title": "overview 37; compact mechanisms 31",
            "section_title": "overview only, fitted to panel width",
            "core_labels": "compact mechanisms 24–29",
            "supporting_labels": "compact mechanisms 21–23",
            "italics": "math variables and selected English condition labels",
            "numbering": "black circles with white numbers placed on execution arrows",
        },
        "diagram_grammar": {
            "entities": ["agent", "FULL page", "Mamba state", "pool", "HiCache", "DMA"],
            "edge_actions": ["D2H", "H2D", "reclaim", "admit"],
            "duration_bars": "arrow-ended timeline intervals, not component nodes",
            "join_candidate": "forecast only; actual child RETURN still required",
            "wait_completion_paths": {
                "child": "async CPU phase/work, runtime rate hint, RCCG-bound parent",
                "tool": "history and elapsed age, conditional end hint, owning agent",
                "output": "candidate identity and timing for cache/admission planning",
            },
            "work_hint": "snapshot remaining work advanced using observed effective rate",
            "data_dependencies": "native layer fences and required state; no global ACK barrier",
            "state_selection": "match the safe-input checkpoint, not simply newest state",
        },
        "palette": PALETTE,
        "math_typography": (
            "editable Liberation Serif italic variables; positioned upright subscripts "
            "and superscripts; editable fraction bars; no rasterized equations"
        ),
        "renderer": "CairoSVG / Pillow / python-pptx",
        "figures": manifest, "powerpoints": powerpoints,
        "standalone_powerpoints": standalone_powerpoints,
    }, indent=2) + "\n")
    print(json.dumps({
        "output": str(args.output), "figures": manifest, "powerpoints": powerpoints,
        "standalone_powerpoints": standalone_powerpoints,
    }, indent=2))


if __name__ == "__main__":
    main()
